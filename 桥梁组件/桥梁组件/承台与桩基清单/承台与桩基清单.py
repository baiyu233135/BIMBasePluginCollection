from pyp3d import *
import os
from openpyxl import *
from openpyxl.styles import *
path = sys.argv[0]
set_global_variable('\a_path', path[0:-3])
entitys = get_element_from_boxselect()

all = set()
temp = []
tempEntity = []
supple_keys = []
for i in entitys:
    if entityid_isvaid(i):
        # 获取模型属性
        datakey = get_datakey_from_entity(i)
        noumenon = get_noumKV_from_instancekey(datakey)
        # "体积检查"
        supple_keys_i = [string for string in list(noumenon.keys()) if '总体积' in string]
        if not supple_keys_i:
            continue
        # 模型属性提资
        d = {}
        for k, j in noumenon.items():
            if not isinstance(j, bytearray):
                d[k] = j
            if ('\x07' not in k) and (k not in Tempfilter):
                all.add(k)
        temp.append(d)
        tempEntity.append(i)
        supple_keys_i.reverse()
        supple_keys.append(supple_keys_i)
if not temp:
    messages("未选择实体！")
    exit()
supplehead = [i+"(mm)" for i in max(supple_keys)]
headName = ["序号","构件名称","构件类别","数量","承台体积/m³","桩基体积/m³"]+supplehead
# 创建表
wb = Workbook()
# 创建sheet
sheet = wb.create_sheet('清单', 0)
# 合并单元格
sheet.merge_cells(start_row=1, start_column=1,
                  end_row=1, end_column=len(headName))
sheet.merge_cells(start_row=2, start_column=1,
                  end_row=2, end_column=len(headName))
sheet.merge_cells(start_row=3, start_column=1,
                  end_row=3, end_column=len(headName))
# 黑边框定义
border = Border(top=Side(border_style='thin', color='FF000000'),
                right=Side(border_style='thin', color='FF000000'),
                bottom=Side(border_style='thin', color='FF000000'),
                left=Side(border_style='thin', color='FF000000'))
# 居中样式
center = Alignment(horizontal='center', vertical='center')
# 字体字号，加粗
font = Font(size=25, bold=True)
# 字体字号，加粗
font2 = Font(size=15, bold=True)
# 填充单元格颜色
fill = PatternFill(fill_type="solid", start_color="6A967E")
# 单元格控制
sheet.cell(row=1, column=1).alignment = center
sheet.cell(row=1, column=1).value = "算量清单"
sheet.cell(row=1, column=1).font = font
sheet.cell(row=2, column=1).value = "工程名称"+'     '*ceil((len(headName)+3)/2)+"工程编号"
sheet.cell(row=2, column=1).font = font2
sheet.cell(row=3, column=1).value = "总包单位"+'     '*ceil((len(headName)+3)/2)+"合同编号"
sheet.cell(row=3, column=1).font = font2
# 表头
for i in range(len(headName)):
    sheet.cell(row=4, column=1+i).alignment = center
    sheet.cell(row=4, column=1+i).value = headName[i]
    sheet.cell(row=4, column=1+i).fill = fill
    sheet.column_dimensions[chr(i+65)].width = len(headName[i])*2.5

# 内容
for i in range(5,len(temp)+5):
    sheet.cell(row=i, column=1).value = i-4
    sheet.cell(row=i, column=2).value = temp[i-5]["构件名称"]
    sheet.cell(row=i, column=3).value = temp[i-5]["构件类型"]
    sheet.cell(row=i, column=4).value = 1
    # 将体积保留两位小数，并设置单元格数字格式为保留两位小数
    vol_ct = temp[i-5].get("承台体积", 0)
    vol_zj = temp[i-5].get("桩基体积", 0)
    try:
        vol_ct_val = round(float(vol_ct), 2)
    except Exception:
        vol_ct_val = vol_ct
    try:
        vol_zj_val = round(float(vol_zj), 2)
    except Exception:
        vol_zj_val = vol_zj
    c5 = sheet.cell(row=i, column=5)
    c5.value = vol_ct_val
    c5.number_format = '0.00'
    c6 = sheet.cell(row=i, column=6)
    c6.value = vol_zj_val
    c6.number_format = '0.00'
    for j,supple_key in enumerate(supple_keys[i-5]):
        val = temp[i-5].get(supple_key, 0)
        try:
            val_num = float(val)
        except Exception:
            val_num = val
        if isinstance(val_num, (int, float)) and val_num > 0:
            cell = sheet.cell(row=i, column=7+j)
            cell.value = round(val_num, 2)
            cell.number_format = '0.00'
# 保存路径
wb.save(os.path.dirname(__file__)+'\\算量清单.xlsx')
wb.close()
os.startfile(os.path.dirname(__file__)+'\\算量清单.xlsx')
