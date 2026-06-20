# -*- coding: utf-8 -*-
from pyp3d import *
import os
import subprocess
import math
import openpyxl  # 新增：用于读取xlsx文件


# 弹出窗口 (保持不变)
def select_main_action(plan_added, vert_added):
    plan_status = "已添加" if plan_added else "未添加"
    vert_status = "已添加" if vert_added else "未添加"
    
    ps_cmd = (
        "Add-Type -AssemblyName System.Windows.Forms; "
        "$form = New-Object System.Windows.Forms.Form; "
        "$form.Text = '要素表批量导入终端'; "
        "$form.Size = New-Object System.Drawing.Size(380,240); "
        "$form.StartPosition = 'CenterScreen'; "
        "$form.TopMost = $true; "
        "$lbl1 = New-Object System.Windows.Forms.Label; "
        "$lbl1.Text = '平面要素表： " + plan_status + "'; "
        "$lbl1.Location = New-Object System.Drawing.Point(20,15); "
        "$lbl1.AutoSize = $true; "
        "$form.Controls.Add($lbl1); "
        "$lbl2 = New-Object System.Windows.Forms.Label; "
        "$lbl2.Text = '竖直要素表： " + vert_status + "'; "
        "$lbl2.Location = New-Object System.Drawing.Point(190,15); "
        "$lbl2.AutoSize = $true; "
        "$form.Controls.Add($lbl2); "
        "$btnP = New-Object System.Windows.Forms.Button; "
        "$btnP.Text = '添加平面要素表'; "
        "$btnP.Location = New-Object System.Drawing.Point(20,50); "
        "$btnP.Size = New-Object System.Drawing.Size(150,40); "
        "$btnP.Add_Click({$global:res='Planar'; $form.Close()}); "
        "$form.Controls.Add($btnP); "
        "$btnV = New-Object System.Windows.Forms.Button; "
        "$btnV.Text = '添加竖直要素表'; "
        "$btnV.Location = New-Object System.Drawing.Point(190,50); "
        "$btnV.Size = New-Object System.Drawing.Size(150,40); "
        "$btnV.Add_Click({$global:res='Vertical'; $form.Close()}); "
        "$form.Controls.Add($btnV); "
        "$btnC = New-Object System.Windows.Forms.Button; "
        "$btnC.Text = '确定并生成 3D 组合模型'; "
        "$btnC.Location = New-Object System.Drawing.Point(20,120); "
        "$btnC.Size = New-Object System.Drawing.Size(320,50); "
        "$btnC.BackColor = [System.Drawing.Color]::LightGreen; "
        "$btnC.Add_Click({$global:res='Confirm'; $form.Close()}); "
        "$form.Controls.Add($btnC); "
        "$form.ShowDialog() | Out-Null; "
        "Write-Host $global:res"
    )
    proc = subprocess.Popen(['powershell', '-Command', ps_cmd], stdout=subprocess.PIPE, shell=True, universal_newlines=True)
    stdout, _ = proc.communicate()
    return stdout.strip()

# 弹出文件对话框 (修改为过滤 xlsx)
def select_excel_file(title):
    default_path = r"C:\ProgramData\PKPM\BIMBase\Plugins\Pro\V1.6\pythonplugin\Sample"
    ps_cmd = (
        "Add-Type -AssemblyName System.Windows.Forms; "
        "$f = New-Object System.Windows.Forms.OpenFileDialog; "
        "$f.Filter = 'Excel文件 (*.xlsx)|*.xlsx'; "
        "$f.InitialDirectory = '" + default_path + "'; "
        "$f.Title = '" + title + "'; "
        "if($f.ShowDialog() -eq 'OK'){ Write-Host $f.FileName }"
    )
    proc = subprocess.Popen(['powershell', '-Command', ps_cmd], stdout=subprocess.PIPE, shell=True, universal_newlines=True)
    stdout, _ = proc.communicate()
    return stdout.strip()

# 工具函数：解析桩号
def parse_station(station_str):
    if station_str is None: return 0.0
    s = str(station_str).strip().upper().replace('K', '')
    if '+' in s:
        parts = s.split('+')
        return float(parts[0]) * 1000 + float(parts[1])
    return float(s) if s else 0.0

# 3. 核心建模类 (保持不变)
class CombinedRouteModel(Component):
    def __init__(self, plan_data, vert_data):
        Component.__init__(self)
        self.plan_data = plan_data
        self.vert_data = vert_data
        self['模型'] = Attr(None, show=True)
        self.replace()

    @export
    def replace(self):
        # A. 解析平面数据 
        plan_segments = []
        curr_s = 0.0
        curr_ang = 0.0  
        curr_x, curr_y = 0.0, 0.0
        
        for row in self.plan_data:
            t = str(row.get('类型', '')).strip().upper()
            if not t: continue
            length = float(row.get('长度', 0)) if row.get('长度') else 0.0
            radius = float(row.get('半径', 0)) if row.get('半径') else 0.0
            
            seg = {
                'start_s': curr_s, 'end_s': curr_s + length,
                'type': t, 'L': length, 'R': radius,
                'x0': curr_x, 'y0': curr_y, 'ang0': curr_ang
            }
            plan_segments.append(seg)
            
            if t == 'Z':
                curr_x += length * math.cos(curr_ang)
                curr_y += length * math.sin(curr_ang)
            elif t == 'Y' and radius != 0:
                d_theta = length / radius
                curr_ang += d_theta
                xc = curr_x - radius * math.sin(seg['ang0'])
                yc = curr_y + radius * math.cos(seg['ang0'])
                curr_x = xc + radius * math.sin(curr_ang)
                curr_y = yc - radius * math.cos(curr_ang)
            curr_s += length
        max_plan_s = curr_s

        # B. 解析竖直数据 
        pvi_list = []
        for row in self.vert_data:
            if '桩号' not in row or '变坡点标高' not in row: continue
            try:
                s = parse_station(row.get('桩号', ''))
                e = float(str(row.get('变坡点标高', '')).strip())
                r = float(str(row.get('竖曲线半径', '')).strip()) if row.get('竖曲线半径') else 0.0
                pvi_list.append({'s': s, 'e': e, 'r': r})
            except: continue
        pvi_list.sort(key=lambda item: item['s'])
        
        # 计算坡度与抛物线控制点
        for i in range(len(pvi_list) - 1):
            p1, p2 = pvi_list[i], pvi_list[i+1]
            slope = (p2['e'] - p1['e']) / (p2['s'] - p1['s']) if p2['s'] != p1['s'] else 0
            p1['slope_out'] = slope
            p2['slope_in'] = slope

        for i in range(1, len(pvi_list) - 1):
            p = pvi_list[i]
            r = p['r']
            if r > 0:
                g1, g2 = p['slope_in'], p['slope_out']
                L = r * abs(g2 - g1)
                T = L / 2
                p['bvc_s'] = p['s'] - T
                p['evc_s'] = p['s'] + T
                p['bvc_e'] = p['e'] - T * g1
            else:
                p['bvc_s'] = p['s']
                p['evc_s'] = p['s']

        # C. 采样函数 
        def get_xy(s):
            for seg in plan_segments:
                if seg['start_s'] <= s <= seg['end_s'] + 1e-5:
                    ds = s - seg['start_s']
                    if seg['type'] == 'Z':
                        return seg['x0'] + ds * math.cos(seg['ang0']), seg['y0'] + ds * math.sin(seg['ang0'])
                    elif seg['type'] == 'Y' and seg['R'] != 0:
                        xc = seg['x0'] - seg['R'] * math.sin(seg['ang0'])
                        yc = seg['y0'] + seg['R'] * math.cos(seg['ang0'])
                        ang_s = seg['ang0'] + ds / seg['R']
                        return xc + seg['R'] * math.sin(ang_s), yc - seg['R'] * math.cos(ang_s)
            return curr_x, curr_y

        def get_z(s):
            if not pvi_list: return 0.0
            if s <= pvi_list[0]['s']: return pvi_list[0]['e']
            if s >= pvi_list[-1]['s']: return pvi_list[-1]['e']
            for i in range(len(pvi_list) - 1):
                p1, p2 = pvi_list[i], pvi_list[i+1]
                if p1['s'] <= s <= p2['s']:
                    if 'bvc_s' in p2 and p2['bvc_s'] < s < p2['evc_s']:
                        g1, g2 = p2['slope_in'], p2['slope_out']
                        L = p2['evc_s'] - p2['bvc_s']
                        x = s - p2['bvc_s']
                        return p2['bvc_e'] + g1 * x + ((g2 - g1) / (2 * L)) * (x ** 2)
                    elif 'bvc_s' in p1 and p1['bvc_s'] < s < p1['evc_s']:
                        g1, g2 = p1['slope_in'], p1['slope_out']
                        L = p1['evc_s'] - p1['bvc_s']
                        x = s - p1['bvc_s']
                        return p1['bvc_e'] + g1 * x + ((g2 - g1) / (2 * L)) * (x ** 2)
                    else:
                        return p1['e'] + p1['slope_out'] * (s - p1['s'])
            return 0.0

        # D. 生成三维空间曲线 
        sample_step = 0.5 
        max_s = min(max_plan_s, pvi_list[-1]['s']) if pvi_list else max_plan_s
        
        segments_3d = []
        prev_x, prev_y = get_xy(0.0)
        prev_z = get_z(0.0)
        
        s = sample_step
        while s <= max_s + sample_step:
            curr_s = min(s, max_s)
            x, y = get_xy(curr_s)
            z = get_z(curr_s)
            
            segments_3d.append(Line(Vec3(prev_x, prev_y, prev_z), Vec3(x, y, z)))
            
            prev_x, prev_y, prev_z = x, y, z
            if curr_s >= max_s: break
            s += sample_step

        if segments_3d:
            self['模型'] = combine(*segments_3d).color(1, 0, 0)

# 4. 执行主逻辑 (修改为 openpyxl 解析)
def run_plugin():
    plan_data = []
    vert_data = []
    
    while True:
        choice = select_main_action(len(plan_data)>0, len(vert_data)>0)
        if choice == 'Confirm':
            if not plan_data or not vert_data: continue
            break
        elif choice in ['Planar', 'Vertical']:
            title = "请选择平面要素表 (.xlsx)" if choice == 'Planar' else "请选择竖直要素表 (.xlsx)"
            file_path = select_excel_file(title)
            if file_path and os.path.exists(file_path):
                try:
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    ws = wb.active
                    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
                    data = []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if not any(row): continue
                        row_dict = {h: v for h, v in zip(headers, row) if h}
                        data.append(row_dict)
                    
                    if choice == 'Planar': plan_data = data
                    else: vert_data = data
                except Exception as e:
                    print(f"读取 Excel 文件失败: {str(e)}")
        else: return

    place(CombinedRouteModel(plan_data, vert_data))

if __name__ == "__main__":
    run_plugin()