# -*- coding: utf-8 -*-
"""
生成《企业证明文书模板.docx》—— 国创赛产业赛道参赛配套文书。

包含三份模板 + 一页使用说明：
  模板一《产品试用证明》（测量/检测公司版，单方盖章）
  模板二《应用前景评价意见》（建筑公司版，单方盖章）
  模板三《产学研合作意向书》（双方盖章）
  附：使用说明

诚实性约束：全部文书为"试用/意向"口径，严禁"已采购/已大规模应用/效益数字"等虚假表述；
公司名、联系人、日期等一律以【 】占位。

格式：A4；正文仿宋四号（14pt）、西文 Times New Roman、1.5 倍行距、首行缩进 2 字符；
模板抬头黑体居中；每份模板独立分页；页脚居中页码。

输出：docs/国创赛ppt/企业证明文书模板.docx
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT_PATH = os.path.join(ROOT, "docs", "国创赛ppt", "企业证明文书模板.docx")

BODY_EAST = "仿宋"
BODY_ASCII = "Times New Roman"
HEAD_EAST = "黑体"


# ---------------------------------------------------------------------------
# 基础工具函数
# ---------------------------------------------------------------------------

def set_font(run, east=BODY_EAST, ascii_f=BODY_ASCII, size=14, bold=False, color=None):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_first_line_indent_chars(paragraph, chars=2, char_size_pt=14):
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(chars * char_size_pt * 20)))


def add_body(doc, text, indent=True, size=14, bold=False, east=BODY_EAST,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, line_spacing=1.5,
             color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.alignment = align
    run = p.add_run(text)
    set_font(run, east=east, size=size, bold=bold, color=color)
    if indent:
        set_first_line_indent_chars(p, 2, size)
    return p


def add_title(doc, text, page_break=False):
    """模板抬头：黑体二号（22pt）加粗居中。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(24)
    pf.space_after = Pt(24)
    pf.line_spacing = 1.5
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    run = p.add_run(text)
    set_font(run, east=HEAD_EAST, ascii_f=HEAD_EAST, size=22, bold=True)
    return p


def add_sub(doc, text):
    """模板内小标题：黑体四号（14pt）加粗，无缩进。"""
    return add_body(doc, text, indent=False, east=HEAD_EAST, bold=True, space_after=6)


def add_right(doc, text, size=14):
    """落款行：右对齐、无缩进。"""
    return add_body(doc, text, indent=False, align=WD_ALIGN_PARAGRAPH.RIGHT,
                    space_after=4, size=size)


def add_small_note(doc, text):
    """模板内注释：仿宋小四（12pt）。"""
    return add_body(doc, text, size=12, space_after=4)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(2.8)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    style = doc.styles["Normal"]
    style.font.name = BODY_ASCII
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_EAST)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    set_font(run, size=9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


# ---------------------------------------------------------------------------
# 模板一：产品试用证明（测量/检测公司版）
# ---------------------------------------------------------------------------

def build_template1(doc):
    add_title(doc, "产品试用证明")
    add_body(doc, "兹有我单位【公司名称】，于【年 月 日】至【年 月 日】组织相关人员对【团队/学校名称】研制的"
                  "“智检桥安——基于人工智能与计算机视觉的桥梁表观损伤智能检测系统”（含 BIMBase 桌面检测插件"
                  "与桥梁数字孪生在线平台）进行了试用。现将试用情况证明如下：")
    add_sub(doc, "一、试用内容")
    add_body(doc, "（1）无人机巡检照片的病害自动识别与分类（裂缝、剥落、露筋、蜂窝麻面、渗水、锈蚀六类）；")
    add_body(doc, "（2）损伤尺寸量化与三维空间定位；")
    add_body(doc, "（3）检测报告自动生成；")
    add_body(doc, "（4）数字孪生平台三维浏览与病害档案查看。")
    add_sub(doc, "二、试用评价")
    add_body(doc, "（以下两档措辞请贵单位按真实试用情况任选其一，删除另一档）")
    add_body(doc, "【积极档】系统运行稳定，功能与描述相符，损伤识别与三维定位功能对提升检测内业效率具有"
                  "积极作用，具备在桥梁检测业务中进一步验证与应用的价值。")
    add_body(doc, "【中性档】系统运行稳定，功能与描述基本相符，试用情况属实。")
    add_body(doc, "特此证明。")
    add_body(doc, "", indent=False)
    add_right(doc, "单位名称（盖章）：【公司名称】")
    add_right(doc, "经办人：【经办人姓名】")
    add_right(doc, "联系电话：【联系电话】")
    add_right(doc, "【年 月 日】")
    add_body(doc, "", indent=False)
    add_small_note(doc, "注：本证明落款日期须晚于真实试用结束日期；试用内容以实际试用情况为准，未试用的功能"
                        "条目请删除。")


# ---------------------------------------------------------------------------
# 模板二：应用前景评价意见（建筑公司版）
# ---------------------------------------------------------------------------

def build_template2(doc):
    add_title(doc, "应用前景评价意见", page_break=True)
    add_body(doc, "我单位【公司名称】对【团队/学校名称】提交的“智检桥安——基于人工智能与计算机视觉的桥梁"
                  "表观损伤智能检测系统”成果资料（功能演示、数字孪生平台、检测报告样例）进行了审阅与评议，"
                  "形成评价意见如下：")
    add_sub(doc, "一、评价要点")
    add_body(doc, "（1）系统面向桥梁运维期表观损伤检测的数字化需求，业务流程设计完整；")
    add_body(doc, "（2）损伤智能识别、三维定位与报告自动生成等功能与检测业务衔接良好；")
    add_body(doc, "（3）系统为辅助决策工具，可作为定期检测与经常性检查的数字化补充手段；")
    add_body(doc, "（4）建议进一步扩大真实工程数据验证范围。")
    add_sub(doc, "二、评价结论")
    add_body(doc, "该系统技术路线合理、成果真实可验，具备在桥梁检测与养护业务中进一步试用与工程化验证的"
                  "应用前景。")
    add_body(doc, "特此评价。")
    add_body(doc, "", indent=False)
    add_right(doc, "评价单位（盖章）：【公司名称】")
    add_right(doc, "评议人：【姓名】")
    add_right(doc, "职务：【职务】")
    add_right(doc, "【年 月 日】")
    add_body(doc, "", indent=False)
    add_small_note(doc, "注：本意见基于成果资料审阅与功能演示作出，不构成采购承诺；评价日期须晚于实际审阅"
                        "与演示日期。")


# ---------------------------------------------------------------------------
# 模板三：产学研合作意向书（双方盖章版）
# ---------------------------------------------------------------------------

def build_template3(doc):
    add_title(doc, "产学研合作意向书", page_break=True)
    add_body(doc, "甲方：【企业名称】", indent=False)
    add_body(doc, "乙方：【团队（学校）名称】", indent=False)
    add_body(doc, "甲乙双方本着“优势互补、产学协同、诚实信用”的原则，就桥梁表观损伤智能检测技术的联合验证"
                  "与成果试用事宜，达成如下合作意向：")
    add_sub(doc, "一、合作内容")
    add_body(doc, "双方围绕乙方研制的“智检桥安——基于人工智能与计算机视觉的桥梁表观损伤智能检测系统”开展"
                  "联合验证与成果试用：（1）甲方提供真实检测业务场景建议与试用反馈；（2）乙方提供系统部署、"
                  "使用培训与技术文档；（3）双方共同总结试用情况，视情况形成试用情况说明。")
    add_sub(doc, "二、甲方职责")
    add_body(doc, "（1）提供业务场景建议，组织相关人员参与系统试用并反馈意见；（2）在数据合规前提下提供"
                  "必要的数据对接指导；（3）对试用中知悉的乙方技术资料负有保密义务。")
    add_sub(doc, "三、乙方职责")
    add_body(doc, "（1）提供系统软件、使用文档与必要的技术支持；（2）根据甲方试用反馈迭代优化系统功能；"
                  "（3）对甲方提供的工程数据依法保密，对外使用时进行脱敏处理。")
    add_sub(doc, "四、知识产权")
    add_body(doc, "（1）乙方既有成果（系统软件、算法流程、技术文档等）的知识产权归乙方所有；（2）合作过程中"
                  "形成的共同成果，其知识产权归属由双方另行书面约定；（3）未经对方书面许可，任何一方不得擅自"
                  "对外转让或许可第三方使用合作成果。")
    add_sub(doc, "五、保密条款")
    add_body(doc, "双方对合作中知悉的对方商业秘密、技术资料与工程数据负有保密义务，保密期限自本意向书签署"
                  "之日起【三】年；乙方因参加竞赛、成果验收等合理用途展示己方成果资料的，不视为违约，但涉及"
                  "甲方工程数据的部分应予脱敏。")
    add_sub(doc, "六、意向声明")
    add_body(doc, "本意向书仅为双方合作意向的确认，不构成任何采购承诺或排他性义务；具体合作内容与商务条款"
                  "由双方另行签订正式协议约定。本意向书一式两份，双方各执一份，自双方盖章之日起生效，有效期"
                  "【一】年。")
    add_body(doc, "", indent=False)
    add_body(doc, "甲方（盖章）：【企业名称】　　　　　　乙方（盖章）：【团队（学校）名称】", indent=False)
    add_body(doc, "代表人：【姓名】　　　　　　　　　　代表人：【姓名】", indent=False)
    add_body(doc, "联系电话：【联系电话】　　　　　　　联系电话：【联系电话】", indent=False)
    add_body(doc, "【年 月 日】　　　　　　　　　　　　【年 月 日】", indent=False)


# ---------------------------------------------------------------------------
# 使用说明页
# ---------------------------------------------------------------------------

def build_instructions(doc):
    add_title(doc, "使用说明", page_break=True)
    add_body(doc, "本文件收录三份企业证明文书模板，用于“智检桥安”项目参加中国国际大学生创新大赛（2026）"
                  "产业赛道时，请合作企业对真实发生的试用行为出具证明。使用前请完整阅读本说明。")

    add_sub(doc, "一、正确顺序：先试用，后盖章")
    add_body(doc, "（1）将数字孪生平台单文件网页（HTML 文件）发送给企业联系人，或携带笔记本电脑到企业现场"
                  "演示桌面检测插件；（2）请对方人员真实操作系统或审阅成果资料，并收集口头或书面反馈；（3）"
                  "根据对方单位类型选择对应模板；（4）填写【 】占位信息后交对方核对；（5）对方确认无误后盖章；"
                  "（6）扫描电子版，放入对策书附录一。")

    add_sub(doc, "二、三份模板适用场景")
    add_body(doc, "（1）模板一《产品试用证明》：适用于测量公司、检测公司——对方实际动手试用过系统后出具；")
    add_body(doc, "（2）模板二《应用前景评价意见》：适用于建筑公司——对方审阅成果资料、观看功能演示后出具"
                  "评价意见，不要求实际操作过系统；")
    add_body(doc, "（3）模板三《产学研合作意向书》：适用于希望体现“校企协同”深度的企业——需双方盖章，"
                  "流程周期较长，建议提前一周以上预约对接。")

    add_sub(doc, "三、红线提醒")
    add_body(doc, "（1）严禁虚构使用规模与效益数字：本套模板全部为“试用/意向”诚实口径，填写时不得改成"
                  "“已采购”“已大规模应用”“带来××万元效益”等表述；")
    add_body(doc, "（2）日期逻辑必须成立：证明落款日期必须晚于真实试用结束日期，试用日期不得早于系统实际"
                  "可演示日期；")
    add_body(doc, "（3）赛事纪律：中国国际大学生创新大赛对弄虚作假实行一票否决，所有证明必须对应真实发生"
                  "的试用或评议行为；")
    add_body(doc, "（4）信息一致：公司名称、联系人、联系电话必须与盖章单位一致，不得借用任何第三方名义。")

    add_sub(doc, "四、盖章与扫描建议")
    add_body(doc, "公章应盖在落款单位名称与日期之上（骑年压月）；扫描分辨率建议 300dpi 以上、彩色扫描；"
                  "插入对策书附录一时保持原比例、清晰可辨，并在附录清单中注明证明类型与出具单位。")


# ---------------------------------------------------------------------------
# 主流程与自检
# ---------------------------------------------------------------------------

def main():
    doc = Document()
    setup_page(doc)
    doc.core_properties.title = "企业证明文书模板"
    doc.core_properties.author = "智检桥安团队"

    build_template1(doc)
    build_template2(doc)
    build_template3(doc)
    build_instructions(doc)

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    doc.save(OUT_PATH)

    # ---------------------------- 自检 ----------------------------
    check = Document(OUT_PATH)
    full_text = "\n".join(p.text for p in check.paragraphs)
    for t in check.tables:
        for row in t.rows:
            for c in row.cells:
                full_text += "\n" + c.text
    parts = ["产品试用证明", "应用前景评价意见", "产学研合作意向书", "使用说明"]
    missing = [p for p in parts if p not in full_text]
    n_placeholder = full_text.count("【")
    body = check.element.body
    from docx.oxml.ns import qn as _qn
    page_breaks = sum(1 for b in body.findall(".//" + _qn("w:br"))
                      if b.get(_qn("w:type")) == "page")
    print("输出文件:", OUT_PATH)
    print("文件大小: %.1f KB" % (os.path.getsize(OUT_PATH) / 1024.0))
    print("段落数:", len(check.paragraphs))
    print("四部分齐全:", "是" if not missing else "否，缺失: %s" % ",".join(missing))
    print("占位符【 】数量: %d" % n_placeholder)
    print("换页符数量（应为3）:", page_breaks)
    print("禁用词核查: 已采购=%d, 大规模应用=%d, 万元=%d（均应仅在红线提醒中出现）"
          % (full_text.count("已采购"), full_text.count("大规模应用"), full_text.count("万元")))


if __name__ == "__main__":
    main()
