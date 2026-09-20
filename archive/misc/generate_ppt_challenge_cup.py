# -*- coding: utf-8 -*-
"""生成 20 页挑战杯产业赛道 PPT：桥元智构"""
import os, glob, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TMP = os.path.join(os.path.dirname(__file__), "_ppt_tmp")
os.makedirs(TMP, exist_ok=True)
OUT = os.path.join(ROOT, "docs", "挑战杯项目官方文件及参考", "挑战杯产业赛道_桥元智构_v1.pptx")

# ---- 主题色 ----
DARK   = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝
BLUE   = RGBColor(0x2E, 0x75, 0xB6)   # 工程蓝
ORANGE = RGBColor(0xED, 0x7D, 0x31)   # 强调橙
GRAY   = RGBColor(0xF2, 0xF2, 0xF2)   # 浅灰
DGRAY  = RGBColor(0x7F, 0x7F, 0x7F)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LBLUE  = RGBColor(0xDE, 0xEA, 0xF6)   # 浅蓝

FONT = "微软雅黑"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

inserted = []   # 记录插入图片
placed = []     # 记录占位框

def _set_font(run, size, color=DARK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    r = run._r.rPr
    ea = r.find(qn('a:ea'))
    if ea is None:
        ea = r.makeelement(qn('a:ea'), {})
        r.append(ea)
    ea.set('typeface', FONT)

def set_text(tf, lines, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT,
             space_after=6, line_spacing=1.0):
    """lines: list of str 或 (text, size, color, bold)"""
    tf.word_wrap = True
    first = True
    for ln in lines:
        if isinstance(ln, tuple):
            text, sz, col, bd = (ln + (None,) * 4)[:4]
            sz = sz or size; col = col or color; bd = bold if bd is None else bd
        else:
            text, sz, col, bd = ln, size, color, bold
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        run = p.add_run(); run.text = text
        _set_font(run, sz, col, bd)

def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, shadow=False):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def add_text(slide, x, y, w, h, lines, **kw):
    tb = slide.shapes.add_textbox(x, y, w, h)
    set_text(tb.text_frame, lines, **kw)
    return tb

def top_bar(slide, title, subtitle=None, page=None):
    add_rect(slide, 0, 0, W, Inches(0.95), fill=DARK)
    add_rect(slide, 0, Inches(0.95), W, Inches(0.06), fill=ORANGE)
    add_text(slide, Inches(0.45), Inches(0.12), Inches(10.2), Inches(0.6),
             [(title, 26, WHITE, True)])
    if subtitle:
        add_text(slide, Inches(0.47), Inches(0.62), Inches(11.5), Inches(0.3),
                 [(subtitle, 11, LBLUE, False)])
    # 右上角项目简称
    add_text(slide, Inches(11.0), Inches(0.28), Inches(2.0), Inches(0.4),
             [("桥元智构", 14, ORANGE, True)], align=PP_ALIGN.RIGHT)
    if page:
        footer(slide, page)

def footer(slide, page):
    add_rect(slide, 0, Inches(7.18), W, Inches(0.32), fill=GRAY)
    add_text(slide, Inches(0.45), Inches(7.2), Inches(5), Inches(0.28),
             [("桥元智构 · 基于国产 BIMBase 与人工智能的桥梁建养一体化平台", 9, DGRAY, False)])
    add_text(slide, Inches(12.3), Inches(7.2), Inches(0.8), Inches(0.28),
             [(f"{page:02d} / 20", 10, DGRAY, True)], align=PP_ALIGN.RIGHT)

def bullets(slide, x, y, w, h, items, size=14, space_after=8):
    """items: list of (level, text) 或 (level, text, color, bold)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        lv, text = it[0], it[1]
        col = it[2] if len(it) > 2 else DARK
        bd = it[3] if len(it) > 3 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lv
        p.space_after = Pt(space_after)
        p.line_spacing = 1.05
        marker = "■ " if lv == 0 else ("▸ " if lv == 1 else "· ")
        run = p.add_run(); run.text = marker + text
        _set_font(run, size if lv == 0 else size - 2, ORANGE if lv == 0 and col is DARK else col, bd if lv == 0 else False)
    return tb

def placeholder(slide, x, y, w, h, text, page):
    sp = add_rect(slide, x, y, w, h, fill=GRAY, line=DGRAY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame
    tf.word_wrap = True
    set_text(tf, [("【截图占位】", 12, DGRAY, True), (text, 11, DGRAY, False)],
             align=PP_ALIGN.CENTER, space_after=4)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    placed.append(f"P{page}: {text}")
    return sp

def img_size(path):
    with Image.open(path) as im:
        return im.size

def add_image_fit(slide, path, x, y, w, h, page, border=True):
    """在 (x,y,w,h) 内等比放置图片并居中，返回实际放置区域"""
    if not os.path.exists(path):
        placeholder(slide, x, y, w, h, os.path.basename(path), page)
        return None
    iw, ih = img_size(path)
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        nw, nh = w, int(w / img_ratio)
    else:
        nh, nw = h, int(h * img_ratio)
    nx = x + int((w - nw) / 2)
    ny = y + int((h - nh) / 2)
    slide.shapes.add_picture(path, nx, ny, nw, nh)
    if border:
        add_rect(slide, nx, ny, nw, nh, fill=None, line=BLUE)
    inserted.append(f"P{page}: {os.path.basename(path)}")
    return nx, ny, nw, nh

def bmp_to_png(pattern, out_prefix):
    files = sorted(glob.glob(os.path.join(ROOT, pattern)))
    outs = []
    for i, f in enumerate(files):
        out = os.path.join(TMP, f"{out_prefix}_{i:02d}.png")
        with Image.open(f) as im:
            im.save(out)
        outs.append((os.path.basename(os.path.dirname(os.path.dirname(f))),
                     os.path.splitext(os.path.basename(f))[0].replace("-缩略图", ""), out))
    return outs

# ================= P1 封面 =================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, fill=DARK)
add_rect(s, 0, Inches(6.9), W, Inches(0.6), fill=BLUE)
add_rect(s, Inches(0.55), Inches(1.7), Inches(0.14), Inches(1.9), fill=ORANGE)
add_text(s, Inches(0.95), Inches(1.55), Inches(8.5), Inches(1.4),
         [("桥元智构", 54, WHITE, True)])
add_text(s, Inches(0.98), Inches(2.85), Inches(8.6), Inches(0.9),
         [("基于国产 BIMBase 三维引擎与人工智能的桥梁建养一体化平台", 20, LBLUE, False)])
add_rect(s, Inches(0.98), Inches(3.75), Inches(7.9), Inches(0.85),
         fill=RGBColor(0x27, 0x5F, 0x8F), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(1.2), Inches(3.88), Inches(7.6), Inches(0.65),
         [("2026“挑战杯”产业赛道 · 揭榜挂帅 XH-202601", 15, ORANGE, True),
          ("人工智能与国产三维图形引擎关键核心技术的融合创新应用研究", 12, WHITE, False)])
add_text(s, Inches(0.98), Inches(5.0), Inches(7.5), Inches(1.4),
         [("参赛学校：________________（待填写）", 14, WHITE, False),
          ("团队名称：________________（待填写）", 14, WHITE, False),
          ("指导教师：________________（待填写）", 14, WHITE, False)], space_after=10)
add_text(s, Inches(0.55), Inches(7.0), Inches(12.2), Inches(0.45),
         [("建 · 验 · 诊 · 养 全链路闭环  |  常泰长江大桥真实工程验证", 13, WHITE, False)],
         align=PP_ALIGN.CENTER)
# 右侧缩略图拼贴
thumbs_all = bmp_to_png("常_常泰长江大桥/常泰长江大桥/*/$*bfa/*-缩略图.bmp", "thumb")
pick = [thumbs_all[i] for i in (0, 2, 5, 8, 11, 13) if i < len(thumbs_all)]
ty = Inches(1.35)
for i, (_, name, p) in enumerate(pick[:6]):
    r, c = divmod(i, 2)
    add_image_fit(s, p, Inches(9.35) + c * Inches(1.95), ty + r * Inches(1.75),
                  Inches(1.85), Inches(1.65), 1, border=False)
add_text(s, Inches(9.35), Inches(6.6), Inches(3.8), Inches(0.3),
         [("常泰长江大桥参数化构件（节选）", 9, LBLUE, False)], align=PP_ALIGN.CENTER)

# ================= P2 目录 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "目录", "CONTENTS", 2)
toc = [
    ("01", "行业背景与痛点", "国产引擎 · 人工巡检 · 数据割裂"),
    ("02", "项目总览与技术底座", "BIMBase 国产引擎 + 四路 AI"),
    ("03", "智能建造 · 建", "AI 建模 · CAD 画板 · 图纸识别"),
    ("04", "参数化组件库", "15 个原创 bfa 参数化构件"),
    ("05", "工程验证 · 常泰长江大桥", "244.9 MB 完整 P3D 工程"),
    ("06", "智能诊断 · 诊", "病害检测 · 大模型诊断 · 面投影定位"),
    ("07", "智慧养护闭环 · 养", "报告自动生成与归档"),
    ("08", "创新点 · 评分对照 · 产业价值", "六项创新 · 满分对照 · 推广路径"),
]
for i, (num, t, d) in enumerate(toc):
    r, c = divmod(i, 4)
    x = Inches(0.6) + c * Inches(6.3)
    y = Inches(1.5) + r * Inches(2.6)
    card = add_rect(s, x, y, Inches(5.9), Inches(2.2), fill=GRAY,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(2.2), fill=BLUE if r == 0 else ORANGE)
    add_text(s, x + Inches(0.35), y + Inches(0.18), Inches(1.4), Inches(0.9),
             [(num, 30, BLUE if r == 0 else ORANGE, True)])
    add_text(s, x + Inches(1.5), y + Inches(0.3), Inches(4.2), Inches(1.7),
             [(t, 17, DARK, True), (d, 11, DGRAY, False)], space_after=8)

# ================= P3 行业背景与痛点 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "行业背景与痛点", "桥梁建养行业的四大结构性难题", 3)
pains = [
    ("引擎受制于人", "桥梁三维设计建模长期依赖国外商业 BIM 软件，国产三维图形引擎存在“卡脖子”风险"),
    ("巡检依赖人工", "桥梁病害检测以人工为主，效率低、主观性强、高空作业危险"),
    ("病害与模型脱节", "检测结果停留在纸质报告，无法落到 BIM 具体构件，难以追踪"),
    ("建养数据割裂", "建造模型与运维养护数据互不贯通，模型价值随竣工而终结"),
]
for i, (t, d) in enumerate(pains):
    r, c = divmod(i, 2)
    x = Inches(0.6) + c * Inches(6.3)
    y = Inches(1.45) + r * Inches(2.45)
    add_rect(s, x, y, Inches(5.9), Inches(2.15), fill=WHITE, line=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(2.15), fill=ORANGE)
    add_text(s, x + Inches(0.35), y + Inches(0.22), Inches(5.2), Inches(1.8),
             [(t, 18, DARK, True), (d, 13, RGBColor(0x40, 0x40, 0x40), False)], space_after=10)
bar = add_rect(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.62), fill=DARK,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(bar.text_frame,
         [("结论：亟需“国产引擎 + 人工智能”的桥梁建养一体化平台", 16, WHITE, True)],
         align=PP_ALIGN.CENTER)
bar.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# ================= P4 项目总览 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "项目总览", "建 — 验 — 诊 — 养 全链路闭环架构", 4)
flow = [
    ("建", "AI 智能建模\nCAD 画板\n图纸识别", BLUE),
    ("参数化\n组件库", "15+ 原创 bfa 构件", DARK),
    ("验", "常泰长江大桥\nP3D 工程", BLUE),
    ("诊", "病害检测\nAI 诊断\n面投影定位", ORANGE),
    ("养", "报告闭环\n归档", BLUE),
]
fx, fy, fw, fh = Inches(0.55), Inches(1.7), Inches(2.25), Inches(2.5)
for i, (t, d, col) in enumerate(flow):
    x = fx + i * Inches(2.5)
    chev = add_rect(s, x, fy, fw, fh, fill=col, shape=MSO_SHAPE.CHEVRON)
    tf = chev.text_frame; tf.word_wrap = True
    set_text(tf, [(t, 22, WHITE, True)] + [(ln, 11, WHITE, False) for ln in d.split("\n")],
             align=PP_ALIGN.CENTER, space_after=4)
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
base = add_rect(s, Inches(0.55), Inches(4.7), Inches(12.2), Inches(0.75), fill=LBLUE,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(base.text_frame,
         [("底座：BIMBase 国产三维图形引擎 · pyp3d 参数化组件 · Python 插件体系", 15, DARK, True)],
         align=PP_ALIGN.CENTER)
base.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
digits = [("9", "大插件"), ("15+", "原创参数化构件"), ("1,000+", "标注样本"), ("1", "条建养全链路闭环")]
for i, (n, d) in enumerate(digits):
    x = Inches(0.55) + i * Inches(3.15)
    add_rect(s, x, Inches(5.85), Inches(2.9), Inches(1.05), fill=WHITE, line=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x, Inches(5.95), Inches(2.9), Inches(0.55),
             [(n, 26, ORANGE, True)], align=PP_ALIGN.CENTER)
    add_text(s, x, Inches(6.5), Inches(2.9), Inches(0.35),
             [(d, 12, DARK, False)], align=PP_ALIGN.CENTER)

# ================= P5 技术底座 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "技术底座", "国产引擎 × 人工智能 × 工程应用的三层技术栈", 5)
layers = [
    ("国产引擎层", "BIMBase（中国建筑科学研究院 / 构力科技）· pyp3d 参数化建模接口：拉伸 / 放样 / 布尔 / 阵列", DARK),
    ("AI 层", "YOLOv8 检测 · OpenCV 图像分析 · Qwen-VL 多模态诊断 · DeepSeek 大模型 · PaddleOCR · PyMuPDF", BLUE),
    ("应用层", "PyQt5 界面 · python-docx 报告 · ODA DWG/DXF 转换 · DXF 图纸生成", ORANGE),
]
for i, (t, d, col) in enumerate(layers):
    y = Inches(1.5) + i * Inches(1.55)
    add_rect(s, Inches(0.55), y, Inches(2.1), Inches(1.3), fill=col,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(0.55), y + Inches(0.42), Inches(2.1), Inches(0.5),
             [(t, 17, WHITE, True)], align=PP_ALIGN.CENTER)
    add_rect(s, Inches(2.85), y, Inches(6.9), Inches(1.3), fill=GRAY,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(3.1), y + Inches(0.22), Inches(6.5), Inches(0.95),
             [(d, 13, DARK, False)])
ix = Inches(10.1)
for i, p in enumerate(["速构智维/Picture/CAD画板.png", "速构智维/Picture/桥隧识别.png"]):
    fp = os.path.join(ROOT, p)
    if os.path.exists(fp):
        add_image_fit(s, fp, ix, Inches(1.5) + i * Inches(1.7), Inches(2.6), Inches(1.5), 5)
cap = add_rect(s, Inches(10.1), Inches(4.95), Inches(2.6), Inches(1.15), fill=LBLUE,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(cap.text_frame, [("自研插件图标", 12, DARK, True), ("CAD 画板 / 桥隧识别", 10, DARK, False)],
         align=PP_ALIGN.CENTER, space_after=4)

# ================= P6 建·AI智能建模 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智能建造 · 建 —— AI 智能建模（AI_Modeling）", "自然语言一句话 → BIMBase 参数化构件", 6)
bullets(s, Inches(0.55), Inches(1.45), Inches(7.6), Inches(5.4), [
    (0, "自然语言一句话 → BIMBase 参数化构件"),
    (0, "双路驱动："),
    (1, "本地中文解析器优先（897 行：汉字数字“五百二百一百”连读切坐标、全半角归一、单位换算、方向词映射）"),
    (1, "DeepSeek 云端大模型兜底"),
    (0, "七级兜底自动放置链（破解商业 BIM API“必须在交互回调链中创建实体”的限制）："),
    (1, "直接几何 → PlaceToDirect → place_to → 烘焙坐标"),
    (1, "→ 临时 IFC/OBJ 导入 → SendInput 模拟交互 → 手动放置"),
    (0, "智能阵列：线性 / 矩形 / 极坐标，沿路线等距布置并自动对齐切向"),
], size=14)
placeholder(s, Inches(8.5), Inches(1.6), Inches(4.25), Inches(3.4), "AI 建模面板运行界面（BIMBase 内运行后替换）", 6)
add_rect(s, Inches(8.5), Inches(5.2), Inches(4.25), Inches(1.3), fill=LBLUE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(8.7), Inches(5.35), Inches(3.9), Inches(1.0),
         [("示例：", 12, DARK, True),
          ("“沿路线每隔 50 米放一个门式桥墩” → 自动阵列布置", 12, DARK, False)], space_after=4)

# ================= P7 建·CAD画板 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智能建造 · 建 —— 速构智维 CAD 画板", "内嵌 BIMBase 的 PyQt 参数化二维画板", 7)
bullets(s, Inches(0.55), Inches(1.4), Inches(7.6), Inches(5.5), [
    (0, "A3 虚拟画布（420×297mm）· 16 个绘图/编辑命令 · 50 步撤销 · 端点/交点/垂足/网格吸附"),
    (0, "画板 ↔ BIMBase 双向同步："),
    (1, "19 个 pyp3d Component 子类 + 四级放置策略"),
    (1, "“从 BIMBase 更新”反向回写（框选策略，避免 1,900+ 实例全场景扫描卡死）"),
    (0, "面编辑创新：三视图拖拽尺寸 → diff 反推参数 → 批量重生成 → 同步实体"),
    (1, "“改图纸即改模型”"),
    (0, "数据互通：DWG↔DXF（ODA File Converter）+ 矢量 PDF（PyMuPDF）导入导出"),
], size=14)
fp = os.path.join(ROOT, "速构智维/Picture/CAD画板.png")
if os.path.exists(fp):
    add_image_fit(s, fp, Inches(8.6), Inches(1.5), Inches(1.5), Inches(1.5), 7, border=False)
placeholder(s, Inches(8.5), Inches(3.2), Inches(4.25), Inches(3.3), "画板运行截图（BIMBase 内运行后替换）", 7)

# ================= P8 建·图纸识别 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智能建造 · 建 —— 图纸识别", "三视图 PDF/DXF 图纸 → 结构化参数 → 参数化构件自动生成", 8)
bullets(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(2.3), [
    (0, "双方案混合：本地规则解析（第一角投影 pdf_view_recognizer）+ Qwen-VL 多模态识别（drawing_recognizer）"),
    (0, "component_matcher 模板匹配 → 生成参数化 .py → 执行建模"),
    (0, "已打通四大构件：索缆锚锭 / 引桥桥墩 / 门式桥墩 / 承台及桩基", ORANGE, True),
], size=14)
add_image_fit(s, os.path.join(ROOT, "组件测试/docs/门式桥墩_图纸预览.png"),
              Inches(0.8), Inches(3.7), Inches(5.7), Inches(3.0), 8)
add_image_fit(s, os.path.join(ROOT, "组件测试/docs/承台及桩基_图纸预览.png"),
              Inches(6.9), Inches(3.7), Inches(5.7), Inches(3.0), 8)

# ================= P9 构件库·常泰 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "参数化组件库 —— 常泰长江大桥", "15 个原创 bfa 参数化构件 · 3 大分类 · 全部 ≥4 个有效可调参数", 9)
groups = {
    "1-上部结构": ("上部结构（7）", BLUE),
    "2-下部结构": ("下部结构（3）", DARK),
    "3-附属设施": ("附属设施（5）", ORANGE),
}
col_x = {"1-上部结构": Inches(0.45), "2-下部结构": Inches(4.85), "3-附属设施": Inches(9.25)}
tw, th = Inches(1.32), Inches(1.18)
for gname, (label, col) in groups.items():
    x0 = col_x[gname]
    add_rect(s, x0, Inches(1.3), Inches(4.3), Inches(0.5), fill=col,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x0, Inches(1.38), Inches(4.3), Inches(0.36),
             [(label, 14, WHITE, True)], align=PP_ALIGN.CENTER)
    items = [t for t in thumbs_all if t[0] == gname]
    for i, (_, name, p) in enumerate(items):
        r, c = divmod(i, 3)
        tx = x0 + c * Inches(1.44)
        ty = Inches(1.95) + r * Inches(1.62)
        add_image_fit(s, p, tx, ty, tw, th, 9, border=False)
        add_text(s, tx, ty + th, tw, Inches(0.3),
                 [(name, 8.5, DARK, False)], align=PP_ALIGN.CENTER)
add_text(s, Inches(0.45), Inches(6.75), Inches(12.4), Inches(0.35),
         [("主塔 · 侧面钢桁架 · 工字钢梁 · 斜拉索 · 支座 · 自密实混凝土层及隔离层 · 轨道底座 / 门式桥墩 · 承台及桩基 · 扩大基础 / 新泽西护栏 · 铁路封闭网 · 轨道 · 接触网 · 路面及交通标线",
           10, DGRAY, False)], align=PP_ALIGN.CENTER)

# ================= P10 构件·组件测试 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "参数化组件库 —— 组件测试（参数化试验场）", "按真实工程图纸逐构件参数化建模与健壮性检验", 10)
bullets(s, Inches(0.55), Inches(1.4), Inches(7.7), Inches(5.5), [
    (0, "4 基础几何体：正方体 / 长方体 / 圆柱 / 直角三棱柱"),
    (0, "5 工程构件（按真实图纸参数化）："),
    (1, "工字钢混凝土组合梁（加劲肋阵列）"),
    (1, "索缆锚锭（异形 Sweep + 14 根底柱）"),
    (1, "引桥桥墩（斜边盖梁 Sweep + 双柱）"),
    (1, "门式桥墩（双变截面空心八边形墩柱，Loft + 布尔减兜底）"),
    (1, "承台及桩基（矩形承台 + 锥形桩阵列）"),
    (0, "工程技巧沉淀：Cone 替代 Cylinder · Combine 合成实体 · mm/cm 单位约定"),
    (0, "合计 20+ 参数化构件，逐个通过参数健壮性检验", ORANGE, True),
], size=13.5)
add_image_fit(s, os.path.join(ROOT, "组件测试/docs/承台及桩基_图纸预览.png"),
              Inches(8.55), Inches(1.5), Inches(4.2), Inches(2.6), 10)
placeholder(s, Inches(8.55), Inches(4.3), Inches(4.2), Inches(2.5), "构件三维效果截图（BIMBase 内运行后替换）", 10)

# ================= P11 验·常泰工程实践 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "工程验证 · 验 —— 常泰长江大桥工程实践", "以 15 个参数化构件在 BIMBase 唯一创作平台搭建完整桥梁", 11)
bullets(s, Inches(0.55), Inches(1.45), Inches(6.9), Inches(5.3), [
    (0, "完整搭建常泰长江大桥：双层斜拉桥，公铁两用"),
    (0, "交付 244.9 MB 完整 P3D 工程文件"),
    (1, "场景自洽、可流畅加载、可直接交付"),
    (0, "工程化能力："),
    (1, "一键 pack_to_bfa 批量打包"),
    (1, "组件库性能优化，独立为 v1.1.0 插件"),
], size=15)
ph = placeholder(s, Inches(7.8), Inches(1.5), Inches(5.0), Inches(5.0),
                 "常泰长江大桥 P3D 全景截图\n（BIMBase 内运行后替换，\n请标注：主塔 / 斜拉索 / 双层桥面 / 门式桥墩）", 11)
add_text(s, Inches(8.0), Inches(6.6), Inches(4.6), Inches(0.35),
         [("部位标注：主塔 · 斜拉索 · 双层桥面 · 门式桥墩", 10, DGRAY, False)], align=PP_ALIGN.CENTER)

# ================= P12 诊·病害检测 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智能诊断 · 诊 —— 病害智能检测", "无人机巡检照片 → 智能检测（6 类病害）", 12)
bullets(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(2.4), [
    (0, "6 类病害：裂缝 / 剥落 / 露筋 / 蜂窝麻面 / 渗水 / 锈蚀"),
    (0, "模式一：YOLOv8n 专用权重，GPU / CPU 均可训练部署"),
    (0, "模式二（创新）：免模型五路 CV 掩码融合 —— 黑帽变换、多尺度自适应阈值、K-Means 颜色聚类、局部纹理标准差、Canny 边缘密度 + HSV 天空抑制 + NMS；无训练数据场景零门槛可用", ORANGE, True),
], size=13.5)
add_image_fit(s, os.path.join(ROOT, "桥梁病害识别/training/runs/models/disease_yolov8n/results.png"),
              Inches(0.7), Inches(4.0), Inches(5.9), Inches(2.9), 12)
add_image_fit(s, os.path.join(ROOT, "桥梁病害识别/runs/detect/val/val_batch0_pred.jpg"),
              Inches(6.9), Inches(4.0), Inches(5.9), Inches(2.9), 12)
add_text(s, Inches(0.7), Inches(3.68), Inches(5.9), Inches(0.3),
         [("YOLOv8n 训练曲线（results.png）", 11, DARK, True)], align=PP_ALIGN.CENTER)
add_text(s, Inches(6.9), Inches(3.68), Inches(5.9), Inches(0.3),
         [("验证集检测效果（val_batch0_pred.jpg）", 11, DARK, True)], align=PP_ALIGN.CENTER)

# ================= P13 诊·数据集与大模型诊断 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智能诊断 · 诊 —— 数据集与大模型诊断", "公开集 + 本地标注 + Qwen-VL 双图结构化诊断", 13)
bullets(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(2.4), [
    (0, "数据：公开集 GYU-DET（11,123 张）+ CODEBRIM（1,022 张）+ 本地巡检报告提取与人工标注（整合集 1,014 张）"),
    (0, "训练：YOLOv8n · 50 epochs · CPU 训练 · mAP50 ≈ 0.336（如实呈现，持续迭代优化）"),
    (0, "Qwen-VL-Max 双图诊断：原图 + 带框图输入 → 严格 JSON 输出（病害类型 / 程度 / 成因分析 / 处置建议 / 紧迫度 五字段）", ORANGE, True),
], size=13.5)
add_image_fit(s, os.path.join(ROOT, "桥梁病害识别/runs/detect/val/BoxPR_curve.png"),
              Inches(0.7), Inches(4.0), Inches(5.9), Inches(2.9), 13)
add_image_fit(s, os.path.join(ROOT, "桥梁病害识别/runs/detect/val/confusion_matrix_normalized.png"),
              Inches(6.9), Inches(4.0), Inches(5.9), Inches(2.9), 13)
add_text(s, Inches(0.7), Inches(3.68), Inches(5.9), Inches(0.3),
         [("精度-召回曲线（BoxPR_curve.png）", 11, DARK, True)], align=PP_ALIGN.CENTER)
add_text(s, Inches(6.9), Inches(3.68), Inches(5.9), Inches(0.3),
         [("归一化混淆矩阵（confusion_matrix_normalized.png）", 11, DARK, True)], align=PP_ALIGN.CENTER)

# ================= P14 诊·面投影 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智能诊断 · 诊 —— 核心创新：2D→3D 解析式面投影", "让病害精准“贴”到 BIM 真实构件", 14)
bullets(s, Inches(0.55), Inches(1.45), Inches(7.7), Inches(5.3), [
    (0, "传统路线：摄影测量 / 点云配准 —— 设备昂贵、流程重、基层难落地"),
    (0, "本项目轻路线："),
    (1, "检测框 → 构件面中心 + 法向 + U/V 轴解析式等比投影"),
    (1, "2mm 法向偏移防闪面，点云标记 Combine 合成"),
    (0, "病害精准“贴”到 BIM 真实构件，构件类型识别覆盖 19+ 种参数签名"),
    (0, "价值：形成“看得见的病害档案”，为养护决策提供空间载体", ORANGE, True),
], size=14)
placeholder(s, Inches(8.55), Inches(1.6), Inches(4.2), Inches(5.0), "病害投影到构件效果截图（BIMBase 内运行后替换）", 14)

# ================= P15 养·闭环 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "智慧养护 · 养 —— 全链路闭环", "从无人机巡检到 Word 报告归档，一键贯通", 15)
steps = ["无人机巡检", "CV / YOLO 检测", "Qwen-VL 诊断", "面投影定位到 BIM 构件", "自动生成 Word 巡检报告"]
sx, sw = Inches(0.4), Inches(2.62)
for i, t in enumerate(steps):
    x = sx + i * Inches(2.56)
    col = [BLUE, BLUE, ORANGE, BLUE, DARK][i]
    ch = add_rect(s, x, Inches(1.5), sw, Inches(1.0), fill=col, shape=MSO_SHAPE.CHEVRON)
    tf = ch.text_frame; tf.word_wrap = True
    set_text(tf, [(t, 12, WHITE, True)], align=PP_ALIGN.CENTER)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
bullets(s, Inches(0.55), Inches(2.9), Inches(7.7), Inches(3.9), [
    (0, "报告（python-docx）：病害清单 / 位置 / 程度 / 成因 / 处置建议，可直接归档"),
    (0, "建-诊-养贯通：建造阶段的模型即养护阶段的数据载体"),
    (1, "模型价值贯穿桥梁全生命周期", ORANGE, True),
], size=14.5)
placeholder(s, Inches(8.55), Inches(2.9), Inches(4.2), Inches(3.6), "自动生成的巡检报告截图（Word 报告，运行后替换）", 15)

# ================= P16 数据参考源 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "数据参考源", "数据按需使用，全程服务于模型训练与识别验证", 16)
cards = [
    ("公开病害数据集", "GYU-DET、CODEBRIM\n（训练与验证）"),
    ("本地巡检数据", "巡检报告提取 + 人工标注\n（整合集 1,014 张）"),
    ("工程图纸与参考资料", "承台及桩基、门式桥墩等\n真实工程图纸"),
    ("数据使用原则", "按需使用，服务于模型训练\n与识别验证，不作他用"),
]
for i, (t, d) in enumerate(cards):
    r, c = divmod(i, 2)
    x = Inches(0.7) + c * Inches(6.3)
    y = Inches(1.7) + r * Inches(2.4)
    add_rect(s, x, y, Inches(5.9), Inches(2.0), fill=GRAY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(2.0), fill=BLUE)
    add_text(s, x + Inches(0.35), y + Inches(0.25), Inches(5.2), Inches(1.6),
             [(t, 17, DARK, True)] + [(ln, 12.5, RGBColor(0x40, 0x40, 0x40), False) for ln in d.split("\n")],
             space_after=8)

# ================= P17 创新点总结 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "创新点总结", "六项核心创新", 17)
innov = [
    ("平台空档", "首个基于国产 BIMBase 的桥梁病害识别—BIM 投影闭环（现有成果均基于 Revit / IFC）"),
    ("免模型检测", "五路 CV 掩码融合，零训练数据门槛"),
    ("多模态诊断", "Qwen-VL 双图输入，五字段结构化输出"),
    ("轻量投影", "解析式 2D→3D 面投影，替代昂贵点云路线"),
    ("混合建模架构", "中文本地解析优先 + 大模型兜底"),
    ("图纸驱动", "面编辑 diff 反推参数，“改图纸即改模型”"),
]
for i, (t, d) in enumerate(innov):
    r, c = divmod(i, 3)
    x = Inches(0.5) + c * Inches(4.25)
    y = Inches(1.5) + r * Inches(2.6)
    add_rect(s, x, y, Inches(4.0), Inches(2.3), fill=WHITE, line=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(4.0), Inches(0.55), fill=DARK if r == 0 else BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.2), y + Inches(0.08), Inches(3.6), Inches(0.4),
             [(f"{i+1}. {t}", 15, WHITE, True)])
    add_text(s, x + Inches(0.25), y + Inches(0.75), Inches(3.5), Inches(1.45),
             [(d, 12, RGBColor(0x40, 0x40, 0x40), False)])

# ================= P18 评分对照 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "评分标准对照", "作品能力逐项对标评分维度", 18)
rows = [
    ("评分维度", "满分", "本项目对应成果"),
    ("作品完整性", "20 分", "244.9 MB P3D 完整工程，场景自洽可交付"),
    ("构件综合完成能力", "25 分", "15+ 原创参数化构件，超额完成、逐个健壮"),
    ("AI 应用能力", "25 分", "检测 + 诊断 + 建模 + 图纸识别四路 AI 深度融合"),
    ("场景及模型美观度", "10 分", "真实重大工程场景：常泰长江大桥双层斜拉桥"),
    ("创新性及应用潜力", "20 分", "建养全链条通用架构，12 项扩展路线可持续演进"),
]
tbl_shape = s.shapes.add_table(6, 3, Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.9))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.9)
tbl.columns[1].width = Inches(1.5)
tbl.columns[2].width = Inches(7.5)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        run = p.add_run(); run.text = val
        if ri == 0:
            _set_font(run, 15, WHITE, True)
            cell.fill.solid(); cell.fill.fore_color.rgb = DARK
        else:
            _set_font(run, 13, DARK, ci == 0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = GRAY if ri % 2 == 0 else WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
bar = add_rect(s, Inches(0.7), Inches(6.65), Inches(11.9), Inches(0.4), fill=None)
add_text(s, Inches(0.7), Inches(6.6), Inches(11.9), Inches(0.4),
         [("合计：100 分 —— 各项指标均给出可核验的工程实证", 13, ORANGE, True)],
         align=PP_ALIGN.CENTER)

# ================= P19 产业价值 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "产业价值与推广", "战略契合 · 生态协同 · 落地路径清晰", 19)
bullets(s, Inches(0.55), Inches(1.45), Inches(12.2), Inches(5.3), [
    (0, "战略契合：交通强国、基础设施智慧运维、国产软件自主可控"),
    (0, "生态协同：与 PKPM 系列软件（结构 / BIM 协同 / CAE）深度结合空间广阔"),
    (0, "推广路径："),
    (1, "设计院（智能建模）→ 施工单位（图纸识别）→ 管养单位（巡检闭环）"),
    (0, "落地门槛低：普通无人机 + 笔记本即可部署，适合基层管养单位"),
    (0, "可持续演进：AI 自动化扩展路线图 12 项（多轮上下文、智能阵列、模板库…）", ORANGE, True),
], size=15, space_after=12)

# ================= P20 结束页 =================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, fill=DARK)
add_rect(s, 0, Inches(4.6), W, Inches(0.05), fill=ORANGE)
add_text(s, Inches(1.0), Inches(1.9), Inches(11.3), Inches(1.1),
         [("让每一座桥梁，建得更快、养得更智", 40, WHITE, True)], align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(3.1), Inches(11.3), Inches(0.6),
         [("以国产引擎筑基，以人工智能赋能 —— 桥元智构，服务桥梁全生命周期", 17, LBLUE, False)],
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.2),
         [("感谢发榜单位与评审专家（待填写）", 14, WHITE, False),
          ("团队联系方式：________________（待填写）", 14, WHITE, False)],
         align=PP_ALIGN.CENTER, space_after=10)
n = len(thumbs_all)
show = [thumbs_all[i] for i in range(n) if i % max(1, n // 8) == 0][:8]
tw2 = Inches(1.15)
total = tw2 * len(show) + Inches(0.15) * (len(show) - 1)
x0 = int((W - total) / 2)
for i, (_, name, p) in enumerate(show):
    add_image_fit(s, p, Emu(int(x0 + i * (tw2 + Inches(0.15)))), Inches(5.95),
                  tw2, Inches(1.0), 20, border=False)

prs.save(OUT)
print("OK ->", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
print("\n--- 插入图片 ---")
for x in inserted:
    print(x)
print("\n--- 占位框 ---")
for x in placed:
    print(x)
