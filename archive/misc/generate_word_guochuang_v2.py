# -*- coding: utf-8 -*-
"""
生成国创赛产业赛道企业命题组《解决方案/命题对策书》Word 文档（v2）。

项目：智检桥安——基于人工智能与计算机视觉的桥梁表观损伤智能检测系统
命题：基于人工智能与计算机视觉的桥梁表观损伤智能检测系统开发
输出：docs/国创赛ppt/智检桥安_命题对策书_v2.docx

v2 相对 v1 的改动：§6.3 数字孪生在线平台整节重写为已完成成果并插入平台截图；
§2.4 部署形态、§9.1 创新点（4）、第十章提交清单第 4 项同步更新口径；其余章节不变。

格式：A4；正文宋体小四（12pt）、1.5 倍行距、段首缩进两字符；
     一级标题微软雅黑三号加粗、二级四号加粗、三级小四加粗。
诚实性约束：mAP50 如实写 0.336、mAP50-95 约 0.085；未实测数据一律标注"待补充实测数据"；
     团队成员以"成员A（待填写）"占位；不编造企业合作名称。
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT_PATH = os.path.join(ROOT, "docs", "国创赛ppt", "智检桥安_命题对策书_v2.docx")
TWIN_IMG_PATH = os.path.join(ROOT, "桥梁病害识别", "数字孪生", "Picture", "平台总览.png")

BODY_EAST = "宋体"
BODY_ASCII = "Times New Roman"
HEAD_EAST = "微软雅黑"
HEADER_FILL = "D9E2F3"  # 表头底色（浅蓝）


# ---------------------------------------------------------------------------
# 基础工具函数
# ---------------------------------------------------------------------------

def set_font(run, east=BODY_EAST, ascii_f=BODY_ASCII, size=12, bold=False, color=None):
    """设置 run 的中英文字体、字号、加粗与颜色。"""
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_first_line_indent_chars(paragraph, chars=2, char_size_pt=12):
    """按字符数设置段首缩进（firstLineChars 优先，firstLine 兜底）。"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(chars * char_size_pt * 20)))


def set_outline_level(paragraph, level):
    """设置大纲级别，便于 Word 导航窗格识别章节。"""
    pPr = paragraph._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level))
    pPr.append(ol)


def add_body(doc, text, indent=True, size=12, bold=False, east=BODY_EAST,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, line_spacing=1.5):
    """正文段落：宋体小四、1.5 倍行距、默认段首缩进两字符。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.alignment = align
    run = p.add_run(text)
    set_font(run, east=east, size=size, bold=bold)
    if indent:
        set_first_line_indent_chars(p, 2, size)
    return p


def add_note(doc, text):
    """注释性文字：宋体五号（10.5pt），段首缩进两字符。"""
    return add_body(doc, text, size=10.5, space_after=4)


def add_h1(doc, text, page_break=True):
    """一级标题：微软雅黑三号（16pt）加粗，默认另起一页。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(14)
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_outline_level(p, 0)
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    run = p.add_run(text)
    set_font(run, east=HEAD_EAST, ascii_f=HEAD_EAST, size=16, bold=True)
    return p


def add_h2(doc, text):
    """二级标题：微软雅黑四号（14pt）加粗。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    set_outline_level(p, 1)
    run = p.add_run(text)
    set_font(run, east=HEAD_EAST, ascii_f=HEAD_EAST, size=14, bold=True)
    return p


def add_h3(doc, text):
    """三级标题：微软雅黑小四（12pt）加粗。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.5
    set_outline_level(p, 2)
    run = p.add_run(text)
    set_font(run, east=HEAD_EAST, ascii_f=HEAD_EAST, size=12, bold=True)
    return p


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=10.5, center=True):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(2)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_font(run, east=BODY_EAST, size=size, bold=bold)


def add_table(doc, caption, headers, rows):
    """带题注的网格表格：表头加粗带底色，单元格宋体五号居中。"""
    cap = add_body(doc, caption, indent=False, bold=True, east=HEAD_EAST,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        set_cell_shading(table.rows[0].cells[i], HEADER_FILL)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, center=(len(str(val)) <= 14))
    add_body(doc, "", indent=False, size=6, space_after=2)
    return cap, table


def add_centered(doc, text, size=14, bold=False, east=BODY_EAST, space_after=6, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    run = p.add_run(text)
    set_font(run, east=east, ascii_f=BODY_ASCII if east == BODY_EAST else east,
             size=size, bold=bold)
    return p


def setup_page(doc):
    """A4 页面、页边距、页眉页脚与 Normal 默认字体。"""
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)
    section.different_first_page_header_footer = True  # 封面不显示页眉页脚

    style = doc.styles["Normal"]
    style.font.name = BODY_ASCII
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_EAST)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = header_p.add_run("智检桥安——基于人工智能与计算机视觉的桥梁表观损伤智能检测系统 · 命题对策书")
    set_font(hr, size=9, color=(0x7F, 0x7F, 0x7F))

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    set_font(run, size=9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


# ---------------------------------------------------------------------------
# 封面与目录
# ---------------------------------------------------------------------------

def build_cover(doc):
    for _ in range(3):
        add_centered(doc, "", size=14)
    add_centered(doc, "中国国际大学生创新大赛（2026）", size=16, bold=True, east=HEAD_EAST)
    add_centered(doc, "产业赛道 · 企业命题组", size=14, bold=True, east=HEAD_EAST)
    add_centered(doc, "", size=14)
    add_centered(doc, "智检桥安", size=36, bold=True, east=HEAD_EAST, space_after=10)
    add_centered(doc, "——基于人工智能与计算机视觉的桥梁表观损伤智能检测系统", size=16, bold=True, east=HEAD_EAST)
    add_centered(doc, "", size=12)
    add_centered(doc, "解决方案 · 命题对策书", size=22, bold=True, east=HEAD_EAST, space_before=6)
    add_centered(doc, "", size=12)
    add_centered(doc, "应答命题：基于人工智能与计算机视觉的桥梁表观损伤智能检测系统开发", size=13)
    add_centered(doc, "（命题企业：广联达）", size=12)
    for _ in range(3):
        add_centered(doc, "", size=14)
    add_centered(doc, "参赛学校：________________（待填写）", size=14)
    add_centered(doc, "团队名称：________________（待填写）", size=14)
    add_centered(doc, "指导教师：________________（待填写）", size=14)
    add_centered(doc, "", size=14)
    add_centered(doc, "2026年8月", size=14)


TOC_ENTRIES = [
    "第一章  企业需求理解与应用场景分析",
    "第二章  总体解决方案",
    "第三章  损伤智能识别与分类",
    "第四章  损伤几何尺寸量化",
    "第五章  三维空间定位",
    "第六章  数字化成果形式",
    "第七章  案例验证与效果评价",
    "第八章  产教协同创新",
    "第九章  创新性与推广价值",
    "第十章  成果提交清单",
    "第十一章  成果应用边界与合规声明",
    "第十二章  实施计划与成果转化路径",
]


def build_toc(doc):
    doc.add_page_break()
    add_centered(doc, "目  录", size=16, bold=True, east=HEAD_EAST, space_after=18)
    for entry in TOC_ENTRIES:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(8)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_first_line_indent_chars(p, 2)
        run = p.add_run(entry)
        set_font(run, size=12)


# ---------------------------------------------------------------------------
# 第一章 企业需求理解与应用场景分析
# ---------------------------------------------------------------------------

def build_ch1(doc):
    add_h1(doc, "第一章  企业需求理解与应用场景分析")

    add_h2(doc, "1.1  行业背景")
    add_body(doc, "桥梁是公路交通网络的关键节点，其技术状况直接关系到人民群众出行安全与国民经济运行效率。"
                  "截至2025年底，全国公路桥梁已超过100万座，大量桥梁陆续进入运维期，定期检测、经常性检查"
                  "与养护维修任务日益繁重，桥梁检测养护市场持续增长。")
    add_body(doc, "表观损伤（如裂缝、剥落、露筋、蜂窝麻面、渗水、锈蚀等）是桥梁结构健康状况最直观的表征，"
                  "是定期检测中必须采集和记录的基础数据，也是养护决策与加固设计的重要依据。随着无人机等低空"
                  "采集装备在检测单位逐步普及，影像数据获取成本显著降低，行业对“采得回、看得懂、算得准、"
                  "定位得到”的智能化检测手段需求日益迫切。")

    add_h2(doc, "1.2  检测业务流程痛点")
    add_body(doc, "本团队对桥梁检测业务全流程进行了梳理，图像采集、人工判读、尺寸测量、空间定位与报告编制"
                  "各环节相互割裂，主要痛点如下：")
    add_body(doc, "（1）图像采集环节：无人机、桥检车抵近采集的照片数量大、角度杂，照片与构件的对应关系依赖"
                  "人工记录，整理归档工作量大。")
    add_body(doc, "（2）人工判读环节：损伤类型与损伤程度的判读依赖检测人员个人经验，主观性强、标准不一；"
                  "登高与桥检车抵近作业风险高、效率低。")
    add_body(doc, "（3）尺寸测量环节：裂缝长度、宽度等几何尺寸需现场人工量测，高空作业危险，测量效率低，"
                  "数据离散度大。")
    add_body(doc, "（4）空间定位环节：病害位置多以“第几跨、第几号墩”等文字方式描述，难以在三维空间中准确"
                  "归档与直观展示，复查时难以快速回溯。")
    add_body(doc, "（5）报告编制环节：采集、判读、测量、定位、报告由不同人员分步完成，数据在多个环节间反复"
                  "誊录，易出错、周期长；检测数据与养护决策、加固设计之间缺乏有效衔接，形成数据断链。")
    add_body(doc, "综上，企业的核心诉求是打通“图像采集→损伤识别→量化分析→空间定位→报告输出”的数字化闭环，"
                  "以数字化手段实现检测全流程的提质增效。")

    add_h2(doc, "1.3  应用场景")
    add_body(doc, "（1）定期检测辅助：为依法开展的定期检测提供影像智能判读结果与检测报告初稿，压缩内业整理时间。")
    add_body(doc, "（2）经常性检查：面向高频次、低成本的日常巡查影像进行批量处理，及时发现病害及其发展趋势。")
    add_body(doc, "（3）灾后应急排查：台风、洪水、地震等灾害发生后，对批量巡检影像进行快速初筛，辅助确定"
                  "重点复查部位。")
    add_body(doc, "（4）无人机航线化巡检：配合航线规划采集的标准化影像，实现照片批量自动分析与病害自动归档。")

    add_h2(doc, "1.4  服务对象")
    add_body(doc, "（1）桥梁检测单位：提升影像判读、尺寸量化与报告编制效率，降低外业作业风险；")
    add_body(doc, "（2）桥梁管养单位：建立可视化的病害数字化档案，支撑养护决策与多期对比；")
    add_body(doc, "（3）设计施工单位：获取带空间定位的病害数据，为加固设计与维修施工提供数据支撑。")

    add_h2(doc, "1.5  系统定位")
    add_body(doc, "本系统定位为桥梁表观损伤智能检测的数字化辅助决策工具，不替代依法开展的定期检测、特殊检测、"
                  "荷载试验和结构安全评估。系统输出的识别、量化与定位结果，均需经具有相应资质的专业人员复核"
                  "确认后，方可作为正式检测成果使用。")


# ---------------------------------------------------------------------------
# 第二章 总体解决方案
# ---------------------------------------------------------------------------

def build_ch2(doc):
    add_h1(doc, "第二章  总体解决方案")

    add_h2(doc, "2.1  研究对象与损伤类型")
    add_body(doc, "本方案选择斜拉桥作为研究桥型，以常泰长江大桥（双层公铁两用斜拉桥）为研究验证对象。选择理由"
                  "如下：其一，斜拉桥构件类型丰富，覆盖主塔、钢桁梁、桥墩、斜拉索、支座等，表观损伤类型具有"
                  "代表性；其二，本团队已建立该桥的BIM参数化模型（P3D格式，244.9MB，含15类参数化构件），具备"
                  "开展三维空间定位与案例验证的现实条件；其三，以真实重大工程为验证对象，能够充分检验方案的"
                  "工程适用性。")
    add_body(doc, "方案针对的表观损伤类型共6类：裂缝、剥落、露筋、蜂窝麻面、渗水、锈蚀。选择理由：该6类损伤"
                  "是混凝土与钢结构桥梁表观检测中出现频率最高、对结构耐久性影响最直接的典型病害，适合通过"
                  "可见光影像进行识别，且与GYU-DET、CODEBRIM等公开数据集的类别体系兼容，便于模型训练与对标。")

    add_h2(doc, "2.2  数字化业务流程")
    add_body(doc, "本方案将检测业务重构为“图像采集→损伤智能识别与分类→损伤几何尺寸量化→三维空间定位→"
                  "检测报告自动生成”五环闭环流程，各环节方法说明如下。")
    add_h3(doc, "2.2.1  图像采集")
    add_body(doc, "以无人机巡检照片为主要数据来源，兼容手机、相机等拍摄影像；模型训练数据来自公开数据集与"
                  "本地自建标注数据集。采集环节对接既有无人机航线化巡检作业模式，不改变检测单位外业习惯。")
    add_h3(doc, "2.2.2  损伤智能识别与分类")
    add_body(doc, "采用双模式检测策略：模式一为YOLOv8n专用检测模型，输出损伤类别与检测框；模式二为免模型"
                  "五路CV掩码融合方案，在无标注数据场景下零门槛冷启动。两种模式互为补充，详见第三章。")
    add_h3(doc, "2.2.3  损伤几何尺寸量化")
    add_body(doc, "基于构件面投影标定获得成像面的物理尺度，将图像中的像素度量换算为物理尺寸，输出损伤长度、"
                  "最大宽度与面积，详见第四章。")
    add_h3(doc, "2.2.4  三维空间定位")
    add_body(doc, "采用解析式2D→3D面投影方法，将图像中的损伤位置映射至BIM构件表面，实现病害与具体构件的"
                  "空间关联和三维可视化标记，详见第五章。")
    add_h3(doc, "2.2.5  检测报告自动生成")
    add_body(doc, "集成Qwen-VL多模态大模型对损伤进行成因诊断，基于python-docx自动生成规范Word检测报告，"
                  "输出病害清单、位置、程度、成因与处置建议，详见第六章。")

    add_h2(doc, "2.3  总体技术架构")
    add_body(doc, "系统采用“数据层—算法层—引擎层—应用层”四层架构，自下而上逐层支撑，如表2-1所示。")
    add_table(doc, "表2-1  总体技术架构",
              ["层次", "主要组成", "功能定位"],
              [["应用层", "BIMBase桌面检测插件、三维可视化展示、Word报告自动生成、网页端数字孪生平台",
                "面向用户的数字化成果形态"],
               ["引擎层", "BIMBase国产三维图形引擎、pyp3d参数化建模库", "三维显示、面投影计算、模型承载"],
               ["算法层", "YOLOv8目标检测（PyTorch）、免模型五路CV掩码融合、Qwen-VL多模态诊断",
                "损伤识别、量化与诊断核心算法"],
               ["数据层", "GYU-DET（11,123张）、CODEBRIM（1,022张）、本地标注整合集（1,014张）",
                "训练验证数据与病害档案数据"]])
    add_body(doc, "四层架构中，数据层为算法层提供训练与验证基础；算法层输出识别、量化与诊断结果；引擎层承载"
                  "三维模型与空间计算；应用层面向最终用户提供可交互的数字化成果。")

    add_h2(doc, "2.4  部署形态")
    add_body(doc, "系统采用“桌面插件＋在线平台”双形态交付。")
    add_body(doc, "（1）桌面检测插件：基于BIMBase平台开发的Python插件，普通笔记本电脑即可运行；YOLOv8n轻量化"
                  "模型支持CPU端推理，无需专业GPU设备，契合基层检测单位硬件条件。")
    add_body(doc, "（2）数字孪生在线平台：团队已建成的网页端桥梁数字孪生平台（基于three.js引擎的单文件离线"
                  "网页），已实现常泰长江大桥真实尺度程序化三维重建与病害三维贴片可视化，浏览器免安装直接"
                  "访问，并支持部署至公网网址对外提供服务，详见第六章。")


# ---------------------------------------------------------------------------
# 第三章 损伤智能识别与分类
# ---------------------------------------------------------------------------

def build_ch3(doc):
    add_h1(doc, "第三章  损伤智能识别与分类")

    add_h2(doc, "3.1  技术路线概述")
    add_body(doc, "损伤智能识别与分类是五环业务流程的核心环节。本方案采用“专用模型＋免模型融合”双模式技术"
                  "路线：以YOLOv8n目标检测模型提供稳定的类别识别能力，以免模型五路CV掩码融合方案解决冷启动"
                  "与无标注场景问题，两者互为补充、覆盖不同应用条件。")

    add_h2(doc, "3.2  YOLOv8n专用检测模型")
    add_body(doc, "模型基于PyTorch深度学习框架构建，选用YOLOv8n轻量化骨干网络，参数量小、推理速度快，支持在"
                  "CPU端完成训练与部署，便于在基层检测单位普通办公电脑上运行。模型覆盖裂缝、剥落、露筋、"
                  "蜂窝麻面、渗水、锈蚀6类表观损伤，训练数据由公开数据集与本地自建标注数据集联合构成，基线"
                  "模型已在CPU环境下完成50个epochs训练。")

    add_h2(doc, "3.3  免模型五路CV掩码融合方案")
    add_body(doc, "针对冷启动场景（无标注数据、无训练条件），本方案设计了免训练的五路CV掩码融合检测方法："
                  "分别通过黑帽变换、多尺度自适应阈值、K-Means颜色聚类、局部纹理标准差、Canny边缘密度五路"
                  "独立通道生成损伤候选掩码，再经HSV天空抑制剔除无关背景、NMS去重合并重叠候选，最终融合输出"
                  "损伤区域。")
    add_body(doc, "该方案无需任何标注数据与训练过程，在新桥型、新场景下可实现零门槛冷启动；其输出结果还可"
                  "作为YOLO模型的预标注辅助，降低数据集建设成本。五路通道从不同图像特征视角刻画损伤，融合后"
                  "对光照变化、表面污渍等干扰具有较好的鲁棒性。")

    add_h2(doc, "3.4  数据集建设")
    add_body(doc, "数据集建设情况如表3-1所示。")
    add_table(doc, "表3-1  数据集建设情况",
              ["数据来源", "规模", "标注方式", "用途"],
              [["GYU-DET（公开数据集）", "11,123张", "已有标注（格式转换）", "训练/验证"],
               ["CODEBRIM（公开数据集）", "1,022张", "已有标注（格式转换）", "训练/验证"],
               ["本地巡检报告提取与人工标注", "1,014张（整合集）", "YOLO格式人工标注",
                "贴近真实业务场景的训练/验证"]])
    add_body(doc, "公开数据集提供类别多样性与规模基础；本地标注数据来自真实巡检报告影像的提取与人工标注，"
                  "使模型更贴近实际业务场景的成像条件与病害形态。")

    add_h2(doc, "3.5  模型性能与提升路线")
    add_body(doc, "当前基线模型（YOLOv8n、50 epochs、CPU训练）的实测性能为：mAP50约0.336，mAP50-95约0.085。"
                  "识别准确率与召回率由PR曲线与归一化混淆矩阵佐证，训练过程曲线完整保存、可复核。上述指标为"
                  "基线模型的真实水平，本对策书如实呈现，不作夸大。")
    add_body(doc, "针对当前基线指标，拟从以下路径持续提升：（1）数据增强：采用旋转、翻转、光照扰动、马赛克"
                  "拼接等策略扩充训练样本多样性；（2）更大骨干网络：升级至YOLOv8s/m等更大骨干并采用迁移学习；"
                  "（3）半监督学习：利用大量未标注巡检影像进行伪标签迭代训练；（4）补充标注与类别均衡采样："
                  "扩充本地标注规模，缓解类别不均衡问题；（5）小目标优化：针对微裂缝等小目标损伤优化特征层"
                  "配置。后续训练完成后，将如实记录并更新各项指标。")


# ---------------------------------------------------------------------------
# 第四章 损伤几何尺寸量化
# ---------------------------------------------------------------------------

def build_ch4(doc):
    add_h1(doc, "第四章  损伤几何尺寸量化")

    add_h2(doc, "4.1  像素—物理尺度换算方法")
    add_body(doc, "损伤几何尺寸量化的核心是将图像中的像素度量换算为物理尺寸。本方案利用2D→3D面投影标定关系"
                  "获得成像面的物理尺度：面投影过程在图像平面与BIM构件表面之间建立了解析式的几何对应关系，"
                  "由此可确定成像面上每个像素对应的实际物理尺寸，进而支撑长度、宽度、面积等几何量的换算。"
                  "该方法无需在检测现场额外布设标尺或标定物，与三维空间定位环节共用同一套几何基础，流程简洁、"
                  "一致性好。")

    add_h2(doc, "4.2  尺寸计算流程")
    add_body(doc, "尺寸量化按“检测框→掩膜细化→骨架提取→量化输出”四步流程处理：")
    add_body(doc, "（1）检测框定位：由识别环节输出损伤所在的矩形检测框，确定量化处理的感兴趣区域；")
    add_body(doc, "（2）掩膜细化：在检测框内对损伤区域进行分割细化，获得贴合损伤真实轮廓的掩膜，剔除背景干扰；")
    add_body(doc, "（3）骨架提取：对裂缝类线状损伤提取中轴骨架，作为长度计算的基准线；")
    add_body(doc, "（4）量化输出：损伤长度取骨架长度乘以像素物理尺寸；最大宽度沿骨架法向逐点扫描取最大值；"
                  "损伤面积取掩膜像素数乘以像素物理尺寸的平方。最终输出损伤的长度、最大宽度与面积三项几何量。")

    add_h2(doc, "4.3  精度验证方法")
    add_body(doc, "精度验证采用与人工测量结果对比的方案：选取覆盖各损伤类型的典型样本，由检测人员按现行检测"
                  "规程进行现场人工测量，作为基准值；系统对同一损伤进行量化，计算系统量化值与人工测量值的"
                  "差异及相对误差，评价量化精度。验证记录表格式如表4-1所示。")
    add_table(doc, "表4-1  尺寸量化精度验证表（测试完成后回填）",
              ["验证项", "人工测量", "系统量化", "差异"],
              [["裂缝长度", "待补充实测数据", "待补充实测数据", "待补充实测数据"],
               ["损伤最大宽度", "待补充实测数据", "待补充实测数据", "待补充实测数据"],
               ["损伤面积", "待补充实测数据", "待补充实测数据", "待补充实测数据"]])
    add_body(doc, "当前各项精度验证尚未完成实测，统一标注为“待补充实测数据”；测试完成后将如实回填，测试脚本"
                  "与原始测量记录同步归档、可备查。")


# ---------------------------------------------------------------------------
# 第五章 三维空间定位
# ---------------------------------------------------------------------------

def build_ch5(doc):
    add_h1(doc, "第五章  三维空间定位")

    add_h2(doc, "5.1  解析式面投影方法原理")
    add_body(doc, "三维空间定位采用解析式2D→3D面投影方法：对BIM模型中的构件面，以面中心为原点、面法向与外"
                  "轮廓U/V轴建立局部坐标系，通过解析公式将图像中的损伤像素坐标等比投影到构件表面，实现损伤"
                  "从二维图像到三维模型的空间映射。投影标记沿构件面法向偏移2mm，避免与构件表面渲染重叠产生"
                  "闪面；全部投影标记以点云形式经Combine合成，在三维模型中持久化展示。")
    add_body(doc, "与依赖密集匹配的三维重建路线不同，解析式面投影直接利用BIM模型已有的精确几何信息，计算量"
                  "小、结果确定性强，且投影过程参数化、可解释、可复核。")

    add_h2(doc, "5.2  损伤—构件空间对应关系")
    add_body(doc, "系统内置构件类型识别能力，覆盖19种以上构件参数签名，涵盖主塔、钢桁架、工字钢梁、斜拉索、"
                  "支座、门式桥墩、承台桩基、护栏、轨道、接触网等常泰长江大桥BIM模型中的主要构件类型。每处"
                  "病害在投影定位后自动归属到具体构件，与构件清单联动，形成“构件—病害”对应的数字化档案，"
                  "支持按构件检索病害、按病害跳转定位，为复查与养护提供空间索引。")

    add_h2(doc, "5.3  与其他定位方法的对比")
    add_body(doc, "本方案与常见定位路线的对比如表5-1所示。")
    add_table(doc, "表5-1  三维空间定位方法对比",
              ["定位方法", "硬件成本", "精度水平", "部署门槛", "适用场景"],
              [["本方案（解析式面投影）", "无额外硬件", "以构件为单元", "低，普通笔记本即可",
                "基层管养单位常态化检测"],
               ["SFM三维重建", "需大量影像与较高算力", "较高", "高，重建耗时较长",
                "科研与重点工程精细分析"],
               ["GPS定位", "依赖差分设备与信号条件", "普通GPS为米级", "中", "开阔场景粗略定位"],
               ["人工参考点测量", "全站仪等测量设备", "高", "中，需专业测量人员", "局部精细测量"]])
    add_body(doc, "对比可见，解析式面投影路线无需额外硬件、计算量小、以构件为定位单元，与检测业务“病害归属"
                  "构件”的管理颗粒度天然匹配，轻量路线更适合基层管养单位推广使用。")

    add_h2(doc, "5.4  定位精度说明")
    add_body(doc, "定位精度的量化测试方案为：在BIM模型中预设若干已知坐标的标记点，经系统投影后对比投影坐标"
                  "与真实坐标的偏差，统计定位误差分布。当前定位精度指标为“待补充实测数据”，测试完成后将如实"
                  "回填，原始记录可备查。")


# ---------------------------------------------------------------------------
# 第六章 数字化成果形式
# ---------------------------------------------------------------------------

def build_ch6(doc):
    add_h1(doc, "第六章  数字化成果形式")
    add_body(doc, "本项目已形成四类可演示、可交互的数字化成果，覆盖检测业务全流程。")

    add_h2(doc, "6.1  BIMBase桌面检测插件")
    add_body(doc, "桌面检测插件基于国产BIMBase三维图形平台开发，实现“无人机巡检照片导入→损伤智能检测→"
                  "Qwen-VL多模态诊断→解析式面投影定位→Word检测报告生成”的全流程一键贯通。插件提供照片导入、"
                  "双模式检测（YOLOv8n与五路CV融合）、结果浏览与三维标记查看等交互功能，是五环业务流程的主要"
                  "承载工具。")

    add_h2(doc, "6.2  检测报告自动生成工具")
    add_body(doc, "报告生成工具基于python-docx库开发，一键生成规范的Word检测报告。报告章节包括：封面信息"
                  "（工程名称、检测日期、检测单位、委托信息）、病害清单（类型、位置（构件）、尺寸、程度、"
                  "紧迫度分级）、损伤位置分布、成因分析与处置建议（按病害类型与程度分级给出）。")
    add_body(doc, "诊断环节采用“原图＋带框图”双图输入Qwen-VL-Max多模态大模型，严格输出包含病害类型、程度、"
                  "成因、处置建议、紧迫度五个字段的结构化JSON，保证报告内容规范、可归档，直接对接检测业务流程。")

    add_h2(doc, "6.3  数字孪生在线平台")
    add_body(doc, "团队已建成常泰长江大桥数字孪生平台，平台为662KB的单文件离线网页，基于three.js三维引擎"
                  "开发，无需安装任何软件，浏览器打开即可使用，同时支持部署至公网网址，实现免安装直接访问。")
    add_body(doc, "平台主要功能包括：（1）真实尺度程序化三维重建：按设计资料对常泰长江大桥进行程序化建模，"
                  "主航道桥为主跨1176m的双塔斜拉桥，孔跨布置142+490+1176+490+142m，空间钻石型主塔总高352m，"
                  "公铁双层钢桁梁桥面；（2）构件显隐控制：15组构件复选框与团队参数化组件库一一对应，支持主塔、"
                  "钢桁梁、斜拉索、支座、门式桥墩、承台及桩基等构件的分组显隐；（3）病害数据接入与三维贴片"
                  "可视化：支持病害数据JSON文件导入，按构件关键词（塔/墩/承台/索/桥面）自动路由落位，在三维"
                  "模型上以贴片形式展示病害位置，点击贴片可查看病害类型、程度、成因与处置建议等详情；（4）"
                  "内置8条示例病害数据，便于无数据环境下演示；（5）照片覆盖尺度换算滑块，支持巡检影像与模型"
                  "尺度的对照调整；（6）工程概况卡集中展示桥梁全长、跨径布置、桥型、主塔、主梁等关键参数。")
    add_body(doc, "平台与桌面检测插件形成“内业处理＋成果展示”的双形态配合：插件完成检测、诊断、定位与报告"
                  "生成，平台承载桥梁三维可视化展示与病害档案服务，二者通过病害数据JSON文件衔接。")
    pic_p = doc.add_paragraph()
    pic_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_after = Pt(2)
    pic_p.add_run().add_picture(TWIN_IMG_PATH, width=Cm(15))
    add_body(doc, "图 6-1  数字孪生平台界面（常泰长江大桥全桥总览）", indent=False, size=9,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    add_h2(doc, "6.4  三维可视化展示")
    add_body(doc, "依托BIMBase引擎的三维显示能力，病害以空间标记形式展示于桥梁BIM模型对应构件表面，支持"
                  "缩放、旋转、漫游等交互浏览；病害标记按紧迫度分级着色，直观呈现全桥病害分布态势，辅助管养"
                  "人员快速掌握桥梁技术状况。")


# ---------------------------------------------------------------------------
# 第七章 案例验证与效果评价
# ---------------------------------------------------------------------------

def build_ch7(doc):
    add_h1(doc, "第七章  案例验证与效果评价")

    add_h2(doc, "7.1  常泰长江大桥验证案例")
    add_body(doc, "本方案以常泰长江大桥为对象开展真实工程验证。该桥为双层公铁两用斜拉桥，团队已建立其BIM"
                  "参数化模型，P3D工程文件244.9MB，包含主塔、钢桁架、工字钢梁、斜拉索、支座、门式桥墩、"
                  "承台桩基、护栏、轨道、接触网等15类参数化构件。")
    add_body(doc, "依托该模型，系统完成了从巡检影像导入、损伤识别、尺寸量化、三维定位到报告生成的全流程"
                  "贯通验证，病害标记成功落位至对应构件表面并与构件清单联动，验证了方案在真实重大工程场景下"
                  "的可行性。")

    add_h2(doc, "7.2  效果评价指标")
    add_body(doc, "案例验证的效果评价指标如表7-1所示。已完成的指标如实填写；尚未完成实测的指标明确标注"
                  "“待补充实测数据”，测试完成后如实回填。全部评价数据真实、清晰、可复核，测试脚本与原始记录"
                  "可备查。")
    add_table(doc, "表7-1  案例验证效果评价指标",
              ["评价指标", "结果"],
              [["损伤识别准确率与召回率",
                "PR曲线与归一化混淆矩阵佐证；mAP50约0.336，mAP50-95约0.085（基线模型，详见第三章）"],
               ["单张图像平均检测时间", "待补充实测数据"],
               ["尺寸量化与人工测量差异", "待补充实测数据"],
               ["空间定位精度", "待补充实测数据"],
               ["人工判读环节减少情况",
                "判读、测量、定位、报告编制四个环节由多人分步完成转变为一人一键流转（定性）"],
               ["检测报告生成效率改善", "一键自动生成规范Word报告，替代手工编制（定性）"]])


# ---------------------------------------------------------------------------
# 第八章 产教协同创新
# ---------------------------------------------------------------------------

def build_ch8(doc):
    add_h1(doc, "第八章  产教协同创新")

    add_h2(doc, "8.1  企业—高校—学生三方协同机制")
    add_body(doc, "本项目按照“企业出题、高校协同、学生攻关、成果服务产业”的模式组织实施：命题企业发布真实"
                  "业务命题，并在检测业务流程、数据格式与业务规范方面提供指导；高校教师从土木工程、智能建造、"
                  "计算机视觉等专业方向把关技术路线；学生团队承担需求分析、方案设计、模型训练、系统开发与案例"
                  "验证的具体攻关任务。三方通过协同机制将产业真实需求转化为可落地的技术成果。")

    add_h2(doc, "8.2  跨专业团队分工")
    add_body(doc, "团队按跨专业原则组建，建议涵盖土木工程、计算机、人工智能、电子信息等专业方向。成员分工"
                  "如表8-1所示。")
    add_table(doc, "表8-1  团队成员分工表",
              ["成员", "专业方向", "主要分工", "具体职责"],
              [["成员A（待填写）", "土木工程（待填写）", "需求分析、案例验证",
                "检测业务流程梳理、损伤类型体系确定、工程案例组织与验证"],
               ["成员B（待填写）", "计算机/人工智能（待填写）", "模型训练、数据集建设",
                "数据集整理与标注、YOLOv8n训练与优化、五路CV融合算法实现"],
               ["成员C（待填写）", "计算机/电子信息（待填写）", "系统开发、集成部署",
                "BIMBase插件开发、面投影定位实现、报告生成与平台搭建"]])
    add_note(doc, "注：成员信息与专业方向待团队确认后填写，最终分工以实际报名材料为准。")

    add_h2(doc, "8.3  校企技术交流机制")
    add_body(doc, "项目拟建立常态化校企技术交流机制：围绕命题要求定期开展线上技术交流，就检测业务流程、"
                  "数据规范与成果形式听取命题企业意见；按阶段向企业与指导教师汇报进展，形成会议纪要等过程"
                  "记录并归档。校企交流的具体过程记录与证明材料待后续补充完善，本对策书不作虚构陈述。")


# ---------------------------------------------------------------------------
# 第九章 创新性与推广价值
# ---------------------------------------------------------------------------

def build_ch9(doc):
    add_h1(doc, "第九章  创新性与推广价值")

    add_h2(doc, "9.1  四大核心创新")
    add_body(doc, "（1）免模型五路CV融合冷启动：黑帽变换、多尺度自适应阈值、K-Means颜色聚类、局部纹理标准差、"
                  "Canny边缘密度五路掩码融合，配合HSV天空抑制与NMS去重，在无标注数据场景下实现零样本冷启动"
                  "检测，并可反哺模型预标注。")
    add_body(doc, "（2）Qwen-VL双图多模态诊断：以“原图＋带框图”双图输入Qwen-VL-Max大模型，严格输出病害类型、"
                  "程度、成因、处置建议、紧迫度五字段结构化JSON，将通用多模态大模型规范化为可归档的工程诊断"
                  "能力。")
    add_body(doc, "（3）解析式面投影轻量定位：利用BIM模型精确几何信息，以解析公式完成2D→3D等比投影，替代"
                  "成本高昂的SFM三维重建与点云路线，无需额外硬件，定位颗粒度与构件管理单元一致。")
    add_body(doc, "（4）识别—量化—定位—报告全链路闭环与双形态交付：五环流程在桌面插件内一键贯通；同时已"
                  "建成基于three.js的数字孪生网页平台，以单文件离线网页形式承载常泰长江大桥真实尺度三维模型与"
                  "病害贴片档案，浏览器免安装直接访问、可部署至公网，实现“桌面检测插件＋已建成数字孪生网页平台”"
                  "的双形态交付。")

    add_h2(doc, "9.2  拓展方向")
    add_body(doc, "结合命题鼓励方向与行业需求，后续拓展方向包括：（1）边缘端轻量化部署：面向无人机、机器人"
                  "等移动平台优化模型，实现现场端实时检测；（2）微裂缝等小目标检测优化：提升对早期微细损伤的"
                  "发现能力；（3）多类型损伤同步识别：扩展损伤类别体系，实现多病害一体化识别；（4）损伤时序"
                  "对比与劣化分析：基于多期检测数据追踪病害发展趋势，支撑预防性养护；（5）与桥梁管理系统、"
                  "养护决策系统的数据对接：打通检测数据到养护决策的“最后一公里”。")

    add_h2(doc, "9.3  迁移推广价值")
    add_body(doc, "本方案的技术方法具有良好的可迁移性：识别、量化、定位、报告各环节均不依赖特定桥型，可迁移"
                  "至隧道、房屋建筑等其他基础设施的表观病害检测场景。团队已在隧道衬砌表观病害方向开展数据"
                  "积累与方法验证，为向“桥隧全域”基础设施检测推广奠定了基础。")


# ---------------------------------------------------------------------------
# 第十章 成果提交清单
# ---------------------------------------------------------------------------

def build_ch10(doc):
    add_h1(doc, "第十章  成果提交清单")
    add_body(doc, "对照产业赛道企业命题的答题要求，本项目的成果提交物共9项，逐项对照如表10-1所示。")
    add_table(doc, "表10-1  成果提交物清单对照表",
              ["序号", "提交物要求", "本项目对应成果", "所在章节", "状态"],
              [["1", "命题对策书/项目技术报告", "本文档", "全文", "已完成"],
               ["2", "企业需求分析和总体解决方案", "需求分析与总体解决方案", "第一、二章", "已完成"],
               ["3", "业务流程图或技术路线图", "五环数字化业务流程、四层技术架构", "第二章", "已完成"],
               ["4", "可演示的数字化成果或交互原型",
                "BIMBase桌面检测插件、数字孪生平台HTML网页（常泰长江大桥数字孪生.html）",
                "第六章", "已完成"],
               ["5", "所选桥型与损伤类型的案例说明", "常泰长江大桥（斜拉桥）、6类表观损伤",
                "第二、七章", "已完成"],
               ["6", "损伤识别、量化与定位的分析结果", "识别指标、量化与定位方法及验证方案",
                "第三、四、五章", "已完成（实测数据待补充）"],
               ["7", "案例验证与应用效果说明", "常泰长江大桥案例验证与效果评价",
                "第七章", "已完成（实测数据待补充）"],
               ["8", "团队成员分工及校企协同过程说明", "团队分工表与校企协同机制",
                "第八章", "待补充（成员信息待填写）"],
               ["9", "数据来源、知识产权及授权情况说明", "数据来源说明与合规声明",
                "第三、十一章", "已完成"]])


# ---------------------------------------------------------------------------
# 第十一章 成果应用边界与合规声明
# ---------------------------------------------------------------------------

def build_ch11(doc):
    add_h1(doc, "第十一章  成果应用边界与合规声明")

    add_h2(doc, "11.1  功能定位边界")
    add_body(doc, "本系统为桥梁表观损伤智能检测的数字化辅助决策工具，其输出结果为检测工作提供辅助参考，"
                  "不替代依法开展的定期检测、特殊检测、荷载试验和结构安全评估，不构成对桥梁结构安全的最终"
                  "判定。正式检测成果必须由具有相应资质的专业人员复核确认并签认。")

    add_h2(doc, "11.2  结果可解释、过程可复核")
    add_body(doc, "系统输出的损伤识别结果、尺寸数据和定位信息均保留中间过程数据，具备依据可查、过程可解释"
                  "和人工可复核的特征：识别结果附置信度与检测框可视化；量化过程保留掩膜与骨架中间结果；定位"
                  "结果可在三维模型中回溯查看。模型训练曲线、PR曲线、混淆矩阵及各项测试脚本与原始记录全部"
                  "归档，可备查验。")

    add_h2(doc, "11.3  数据合规与脱敏")
    add_body(doc, "项目在数据采集与使用过程中严格遵守相关法律法规：涉及实际工程图像、检测报告、企业数据及"
                  "人员信息时，依法取得授权或进行脱敏处理后方可使用；对外展示与参赛材料中不出现未授权的"
                  "工程敏感信息。")

    add_h2(doc, "11.4  知识产权声明")
    add_body(doc, "本项目成果真实合法，不侵犯他人知识产权。公开数据集（GYU-DET、CODEBRIM）按其各自许可协议"
                  "规范使用并标注来源；本地巡检影像的提取与标注数据由团队自建，相关权利归团队所有；BIMBase、"
                  "pyp3d、PyTorch、YOLOv8、Qwen-VL等第三方平台与开源组件均按其授权方式合规使用。系统插件、"
                  "算法流程与本对策书文档为团队原创成果。")


# ---------------------------------------------------------------------------
# 第十二章 实施计划与成果转化路径
# ---------------------------------------------------------------------------

def build_ch12(doc):
    add_h1(doc, "第十二章  实施计划与成果转化路径")

    add_h2(doc, "12.1  阶段实施计划（2026年8月—11月）")
    add_body(doc, "项目按四个阶段推进，计划安排如表12-1所示。")
    add_table(doc, "表12-1  阶段实施计划表",
              ["时间", "阶段目标", "主要任务", "阶段交付物"],
              [["2026年8月", "方案完善与数据补充",
                "本地数据补充标注；实测数据采集方案制定；命题对策书编制", "命题对策书（本文档）"],
               ["2026年9月", "模型优化与系统集成",
                "数据增强与更大骨干网络对比实验；半监督训练探索；插件全流程联调", "优化后模型、插件稳定版"],
               ["2026年10月", "案例验证与实测回填",
                "常泰长江大桥全流程验证；量化精度、定位精度、单张耗时等实测并回填", "案例验证报告、实测数据记录"],
               ["2026年11月", "成果固化与提交",
                "对策书修订；报告样例与演示材料制作；提交物整理与自查", "全套参赛提交物"]])

    add_h2(doc, "12.2  成果转化路径")
    add_body(doc, "成果转化按三步走路径推进：第一步，检测单位工具化落地，以桌面插件授权（License）与按项目"
                  "检测服务的形式，在检测单位实现工具化应用；第二步，平台订阅推广，依托数字孪生在线平台提供"
                  "账号订阅、病害档案托管、多期数据对比与劣化趋势分析服务，面向管养单位推广；第三步，桥隧全域"
                  "推广，将方法迁移至隧道、房屋等基础设施检测场景，形成面向基础设施运维的通用解决方案。本方案"
                  "落地门槛低，普通无人机加笔记本电脑即可完成部署，契合基层管养单位的预算与硬件条件。")

    add_h2(doc, "12.3  项目阶段定位")
    add_body(doc, "本项目当前处于小试阶段：核心功能已在常泰长江大桥真实工程BIM模型上完成全流程贯通验证，"
                  "基线模型指标如实可查，双形态成果可演示、可交互。后续将按照阶段计划扩充数据规模、优化模型"
                  "指标、完成实测回填，逐步向中试与应用示范阶段推进。")


# ---------------------------------------------------------------------------
# 主流程与自检
# ---------------------------------------------------------------------------

def main():
    if not os.path.exists(TWIN_IMG_PATH):
        raise FileNotFoundError("平台截图不存在: %s" % TWIN_IMG_PATH)

    doc = Document()
    setup_page(doc)
    doc.core_properties.title = "智检桥安 · 命题对策书（v2）"
    doc.core_properties.author = "智检桥安团队"
    doc.core_properties.subject = "中国国际大学生创新大赛（2026）产业赛道企业命题组"

    build_cover(doc)
    build_toc(doc)
    build_ch1(doc)
    build_ch2(doc)
    build_ch3(doc)
    build_ch4(doc)
    build_ch5(doc)
    build_ch6(doc)
    build_ch7(doc)
    build_ch8(doc)
    build_ch9(doc)
    build_ch10(doc)
    build_ch11(doc)
    build_ch12(doc)

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    doc.save(OUT_PATH)

    # ---------------------------- 自检 ----------------------------
    check = Document(OUT_PATH)
    para_count = len(check.paragraphs)
    table_count = len(check.tables)
    full_text = "\n".join(p.text for p in check.paragraphs)
    for t in check.tables:
        for row in t.rows:
            for c in row.cells:
                full_text += "\n" + c.text
    chapters = ["第一章", "第二章", "第三章", "第四章", "第五章", "第六章",
                "第七章", "第八章", "第九章", "第十章", "第十一章", "第十二章"]
    missing = [c for c in chapters if c not in full_text]
    placeholders_tianxie = full_text.count("待填写")
    placeholders_shice = full_text.count("待补充实测数据")
    file_size = os.path.getsize(OUT_PATH)

    print("输出文件:", OUT_PATH)
    print("文件大小: %.1f KB" % (file_size / 1024.0))
    print("段落数:", para_count)
    print("表格数:", table_count)
    print("嵌入图片数:", len(check.inline_shapes))
    print("章节齐全:", "是" if not missing else "否，缺失: %s" % ",".join(missing))
    print("占位统计: 待填写 x%d, 待补充实测数据 x%d" % (placeholders_tianxie, placeholders_shice))
    print("关键指标核查: mAP50约0.336 ->", "0.336" in full_text,
          "; mAP50-95约0.085 ->", "0.085" in full_text)
    print("v2口径核查: three.js ->", "three.js" in full_text,
          "; 图6-1 ->", "图 6-1" in full_text)


if __name__ == "__main__":
    main()
