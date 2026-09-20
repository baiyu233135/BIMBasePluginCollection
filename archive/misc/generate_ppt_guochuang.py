# -*- coding: utf-8 -*-
"""生成 20 页国创赛产业赛道 PPT：智检桥安"""
import os, glob, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TMP = os.path.join(os.path.dirname(__file__), "_ppt_tmp")
os.makedirs(TMP, exist_ok=True)
OUT = os.path.join(ROOT, "docs", "国创赛ppt", "国创赛产业赛道_智检桥安_v1.pptx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

# ---- 主题色 ----
DARK   = RGBColor(0x1F, 0x4E, 0x79)
BLUE   = RGBColor(0x2E, 0x75, 0xB6)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
GRAY   = RGBColor(0xF2, 0xF2, 0xF2)
DGRAY  = RGBColor(0x7F, 0x7F, 0x7F)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LBLUE  = RGBColor(0xDE, 0xEA, 0xF6)
BODY   = RGBColor(0x40, 0x40, 0x40)

FONT = "微软雅黑"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

inserted = []
placed = []

def _set_font(run, size, color=DARK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    r = run._r.rPr
    ea = r.find(qn('a:ea'))
    if ea is None:
        ea = r.makeelement(qn('a:ea'), {})
        r.append(ea)
    ea.set('typeface', FONT)

def set_text(tf, lines, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT,
             space_after=6, line_spacing=1.0):
    tf.word_wrap = True
    first = True
    for ln in lines:
        if isinstance(ln, tuple):
            text, sz, col, bd = (ln + (None,) * 4)[:4]
            sz = sz or size; col = col or color; bd = bold if bd is None else bd
        else:
            text, sz, col, bd = ln, size, color, bold
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        run = p.add_run(); run.text = text
        _set_font(run, sz, col, bd)

def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, shadow=False):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def add_text(slide, x, y, w, h, lines, **kw):
    tb = slide.shapes.add_textbox(x, y, w, h)
    set_text(tb.text_frame, lines, **kw)
    return tb

def top_bar(slide, title, subtitle=None, page=None):
    add_rect(slide, 0, 0, W, Inches(0.95), fill=DARK)
    add_rect(slide, 0, Inches(0.95), W, Inches(0.06), fill=ORANGE)
    add_text(slide, Inches(0.45), Inches(0.12), Inches(10.2), Inches(0.6),
             [(title, 26, WHITE, True)])
    if subtitle:
        add_text(slide, Inches(0.47), Inches(0.62), Inches(11.5), Inches(0.3),
                 [(subtitle, 11, LBLUE, False)])
    add_text(slide, Inches(11.0), Inches(0.28), Inches(2.0), Inches(0.4),
             [("智检桥安", 14, ORANGE, True)], align=PP_ALIGN.RIGHT)
    if page:
        footer(slide, page)

def footer(slide, page):
    add_rect(slide, 0, Inches(7.18), W, Inches(0.32), fill=GRAY)
    add_text(slide, Inches(0.45), Inches(7.2), Inches(5), Inches(0.28),
             [("智检桥安 · 基于人工智能与计算机视觉的桥梁表观损伤智能检测系统", 9, DGRAY, False)])
    add_text(slide, Inches(12.3), Inches(7.2), Inches(0.8), Inches(0.28),
             [(f"{page:02d} / 20", 10, DGRAY, True)], align=PP_ALIGN.RIGHT)

def bullets(slide, x, y, w, h, items, size=14, space_after=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        lv, text = it[0], it[1]
        col = it[2] if len(it) > 2 else DARK
        bd = it[3] if len(it) > 3 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lv
        p.space_after = Pt(space_after)
        p.line_spacing = 1.05
        marker = "■ " if lv == 0 else ("▸ " if lv == 1 else "· ")
        run = p.add_run(); run.text = marker + text
        _set_font(run, size if lv == 0 else size - 2, ORANGE if lv == 0 and col is DARK else col, bd if lv == 0 else False)
    return tb

def placeholder(slide, x, y, w, h, text, page):
    sp = add_rect(slide, x, y, w, h, fill=GRAY, line=DGRAY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame
    tf.word_wrap = True
    set_text(tf, [("【截图占位】", 12, DGRAY, True), (text, 11, DGRAY, False)],
             align=PP_ALIGN.CENTER, space_after=4)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    placed.append(f"P{page}: {text}")
    return sp

def img_size(path):
    with Image.open(path) as im:
        return im.size

def add_image_fit(slide, path, x, y, w, h, page, border=True):
    if not os.path.exists(path):
        placeholder(slide, x, y, w, h, os.path.basename(path), page)
        return None
    iw, ih = img_size(path)
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        nw, nh = w, int(w / img_ratio)
    else:
        nh, nw = h, int(h * img_ratio)
    nx = x + int((w - nw) / 2)
    ny = y + int((h - nh) / 2)
    slide.shapes.add_picture(path, nx, ny, nw, nh)
    if border:
        add_rect(slide, nx, ny, nw, nh, fill=None, line=BLUE)
    inserted.append(f"P{page}: {os.path.basename(path)}")
    return nx, ny, nw, nh

def bmp_to_png(pattern, out_prefix):
    files = sorted(glob.glob(os.path.join(ROOT, pattern)))
    outs = []
    for i, f in enumerate(files):
        out = os.path.join(TMP, f"{out_prefix}_{i:02d}.png")
        with Image.open(f) as im:
            im.save(out)
        outs.append((os.path.basename(os.path.dirname(os.path.dirname(f))),
                     os.path.splitext(os.path.basename(f))[0].replace("-缩略图", ""), out))
    return outs

def chevron_row(slide, items, x0, y, cw, ch, gap=0.1, colors=None, title_size=13, note_size=9):
    """items: list of (title, note)；note 为环下小字"""
    n = len(items)
    for i, (t, note) in enumerate(items):
        x = x0 + i * Inches((cw + gap))
        col = (colors[i] if colors else BLUE)
        chev = add_rect(slide, x, y, Inches(cw), Inches(ch), fill=col, shape=MSO_SHAPE.CHEVRON)
        tf = chev.text_frame; tf.word_wrap = True
        set_text(tf, [(t, title_size, WHITE, True)], align=PP_ALIGN.CENTER)
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
        if note:
            add_text(slide, x + Inches(0.15), y + Inches(ch + 0.08), Inches(cw - 0.3), Inches(0.75),
                     [(note, note_size, DGRAY, False)], align=PP_ALIGN.CENTER, space_after=2)

def simple_table(slide, rows, x, y, w, h, col_widths, header_fill=DARK, size=13):
    tbl = slide.shapes.add_table(len(rows), len(col_widths), x, y, w, h).table
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = cw
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run(); run.text = val
            if ri == 0:
                _set_font(run, size + 1, WHITE, True)
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            else:
                _set_font(run, size, DARK, ci == 0)
                cell.fill.solid()
                cell.fill.fore_color.rgb = GRAY if ri % 2 == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tbl

# 图片路径
VAL = os.path.join(ROOT, "桥梁病害识别", "runs", "detect", "val")
TRAIN_RUN = os.path.join(ROOT, "桥梁病害识别", "training", "runs", "models", "disease_yolov8n")
def img(*parts):
    return os.path.join(*parts)

thumbs_all = bmp_to_png("常_常泰长江大桥/常泰长江大桥/*/$*bfa/*-缩略图.bmp", "thumb_gc")

# ================= P1 封面 =================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, fill=DARK)
add_rect(s, 0, Inches(6.9), W, Inches(0.6), fill=BLUE)
add_rect(s, Inches(0.55), Inches(1.55), Inches(0.14), Inches(1.85), fill=ORANGE)
add_text(s, Inches(0.95), Inches(1.35), Inches(8.5), Inches(1.4),
         [("智检桥安", 54, WHITE, True)])
add_text(s, Inches(0.98), Inches(2.6), Inches(8.6), Inches(0.9),
         [("基于人工智能与计算机视觉的桥梁表观损伤智能检测系统", 20, LBLUE, False)])
add_rect(s, Inches(0.98), Inches(3.5), Inches(7.9), Inches(1.1),
         fill=RGBColor(0x27, 0x5F, 0x8F), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(1.2), Inches(3.62), Inches(7.6), Inches(0.9),
         [("中国国际大学生创新大赛（2026）· 产业赛道 · 企业命题应答", 15, ORANGE, True),
          ("命题：基于人工智能与计算机视觉的桥梁表观损伤智能检测系统开发（命题企业：广联达）", 12, WHITE, False)],
         space_after=6)
add_text(s, Inches(0.98), Inches(4.85), Inches(7.5), Inches(1.4),
         [("参赛学校：________________（待填写）", 14, WHITE, False),
          ("团队名称：________________（待填写）", 14, WHITE, False),
          ("指导教师：________________（待填写）", 14, WHITE, False)], space_after=10)
add_text(s, Inches(0.55), Inches(7.0), Inches(12.2), Inches(0.45),
         [("识别 · 量化 · 定位 · 报告 全链路闭环  |  常泰长江大桥真实工程验证", 13, WHITE, False)],
         align=PP_ALIGN.CENTER)
pick = [thumbs_all[i] for i in (0, 2, 5, 8, 11, 13) if i < len(thumbs_all)]
ty = Inches(1.35)
for i, (_, name, p) in enumerate(pick[:6]):
    r, c = divmod(i, 2)
    add_image_fit(s, p, Inches(9.35) + c * Inches(1.95), ty + r * Inches(1.75),
                  Inches(1.85), Inches(1.65), 1, border=False)
add_text(s, Inches(9.35), Inches(6.6), Inches(3.8), Inches(0.3),
         [("常泰长江大桥参数化构件（节选）", 9, LBLUE, False)], align=PP_ALIGN.CENTER)

# ================= P2 目录 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "目录", "CONTENTS", 2)
toc = [
    ("01", "需求与场景", "企业需求 · 痛点 · 应用场景与服务对象"),
    ("02", "总体方案", "桥型与损伤类型 · 数字化业务流程 · 技术架构"),
    ("03", "核心能力", "识别 · 量化 · 定位 · 报告"),
    ("04", "案例验证", "常泰长江大桥 · 效果评价"),
    ("05", "产教协同与创新", "三方协同 · 四大创新 · 成果清单 · 合规边界"),
    ("06", "商业前景", "商业模式 · 推广路径"),
]
for i, (num, t, d) in enumerate(toc):
    r, c = divmod(i, 3)
    x = Inches(0.6) + c * Inches(6.3)
    y = Inches(1.6) + r * Inches(2.55)
    add_rect(s, x, y, Inches(5.9), Inches(2.15), fill=GRAY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(2.15), fill=BLUE if r == 0 else ORANGE)
    add_text(s, x + Inches(0.35), y + Inches(0.18), Inches(1.4), Inches(0.9),
             [(num, 30, BLUE if r == 0 else ORANGE, True)])
    add_text(s, x + Inches(1.5), y + Inches(0.35), Inches(4.2), Inches(1.6),
             [(t, 18, DARK, True), (d, 11.5, DGRAY, False)], space_after=10)

# ================= P3 企业需求与痛点 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "企业需求与痛点", "截至 2025 年底全国公路桥梁已超 100 万座，大量桥梁进入运维期，定期检测任务繁重", 3)
pains = [
    ("低效危险", "登高 / 桥检车抵近人工检测，作业风险高、效率低"),
    ("判读主观", "损伤判读依赖个人经验，主观性强、标准不一"),
    ("流程割裂", "采集 / 识别 / 测量 / 定位 / 报告由不同人员分步完成"),
    ("数据断链", "检测结果与养护决策、加固设计缺乏数据衔接"),
]
for i, (t, d) in enumerate(pains):
    r, c = divmod(i, 2)
    x = Inches(0.6) + c * Inches(6.3)
    y = Inches(1.5) + r * Inches(2.25)
    add_rect(s, x, y, Inches(5.9), Inches(2.0), fill=WHITE, line=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(2.0), fill=ORANGE)
    add_text(s, x + Inches(0.35), y + Inches(0.2), Inches(5.2), Inches(1.7),
             [(f"{i+1}. {t}", 18, DARK, True), (d, 13, BODY, False)], space_after=10)
bar = add_rect(s, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.75), fill=DARK,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(bar.text_frame,
         [("企业核心诉求：打通“图像采集 → 损伤识别 → 量化分析 → 空间定位 → 报告输出”数字化闭环", 15, WHITE, True)],
         align=PP_ALIGN.CENTER)
bar.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# ================= P4 应用场景与服务对象 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "应用场景与服务对象", "深入理解企业检测业务流程，衔接各环节", 4)
scenes = [
    ("定期检测辅助", "为定期检测提供影像判读与报告初稿"),
    ("经常性检查", "高频次、低成本的日常巡查影像处理"),
    ("灾后应急排查", "台风 / 洪水 / 地震后快速批量初筛"),
    ("无人机航线化巡检", "航线规划采集，照片批量自动分析"),
]
for i, (t, d) in enumerate(scenes):
    x = Inches(0.6) + i * Inches(3.15)
    add_rect(s, x, Inches(1.5), Inches(2.9), Inches(1.9), fill=GRAY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.5), Inches(2.9), Inches(0.55), fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x, Inches(1.6), Inches(2.9), Inches(0.4),
             [(t, 14, WHITE, True)], align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.2), Inches(2.5), Inches(1.1),
             [(d, 11.5, BODY, False)])
obj = add_rect(s, Inches(0.6), Inches(3.7), Inches(12.1), Inches(0.75), fill=LBLUE,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(obj.text_frame,
         [("服务对象：桥梁检测单位 · 管养单位 · 设计施工单位", 15, DARK, True)],
         align=PP_ALIGN.CENTER)
obj.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
add_rect(s, Inches(0.6), Inches(4.7), Inches(12.1), Inches(1.15), fill=WHITE, line=ORANGE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.85), Inches(4.85), Inches(11.6), Inches(0.9),
         [("定位声明：本系统为桥梁表观损伤智能检测的数字化辅助决策工具，", 14, DARK, True),
          ("不替代依法开展的定期检测、特殊检测、荷载试验和结构安全评估", 14, ORANGE, True)],
         align=PP_ALIGN.CENTER, space_after=4)
add_text(s, Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.8),
         [("流程衔接：深入理解企业检测业务流程，衔接图像采集、人工判读、尺寸测量、空间定位与报告编制各环节", 12.5, BODY, False)],
         align=PP_ALIGN.CENTER)

# ================= P5 桥型与损伤类型 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "桥型与损伤类型", "以真实重大工程为研究验证对象，覆盖六类典型表观损伤", 5)
bullets(s, Inches(0.55), Inches(1.4), Inches(6.6), Inches(2.4), [
    (0, "典型桥型：斜拉桥"),
    (1, "以常泰长江大桥（双层公铁两用斜拉桥）为研究验证对象"),
    (1, "构件覆盖主塔 / 钢桁梁 / 桥墩 / 斜拉索"),
], size=14)
# 6 类损伤卡片
dmg = ["裂缝", "剥落", "露筋", "蜂窝麻面", "渗水", "锈胀"]
for i, t in enumerate(dmg):
    r, c = divmod(i, 3)
    x = Inches(0.55) + c * Inches(2.25)
    y = Inches(3.6) + r * Inches(1.35)
    cd = add_rect(s, x, y, Inches(2.05), Inches(1.1), fill=DARK if r == 0 else BLUE,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    set_text(cd.text_frame, [(t, 16, WHITE, True)], align=PP_ALIGN.CENTER)
    cd.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
add_text(s, Inches(0.55), Inches(6.5), Inches(6.6), Inches(0.4),
         [("损伤类型（6 类）", 12, DGRAY, True)], align=PP_ALIGN.CENTER)
# 右侧缩略图 4 张
pick5 = [t for t in thumbs_all if t[1] in ("主塔", "门式桥墩", "承台及桩基", "斜拉索")]
for i, (_, name, p) in enumerate(pick5[:4]):
    r, c = divmod(i, 2)
    add_image_fit(s, p, Inches(7.6) + c * Inches(2.6), Inches(1.45) + r * Inches(2.1),
                  Inches(2.45), Inches(1.75), 5, border=False)
    add_text(s, Inches(7.6) + c * Inches(2.6), Inches(1.45) + r * Inches(2.1) + Inches(1.75),
             Inches(2.45), Inches(0.3), [(name, 10, DARK, False)], align=PP_ALIGN.CENTER)
add_text(s, Inches(7.6), Inches(5.85), Inches(5.0), Inches(0.35),
         [("常泰长江大桥参数化构件缩略图", 10, DGRAY, False)], align=PP_ALIGN.CENTER)

# ================= P6 数字化业务流程 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "数字化业务流程", "五环闭环：从图像到报告一键贯通", 6)
flow = [
    ("图像采集", "公开数据集 + 自建数据集 + 现场巡检图像"),
    ("损伤智能识别\n与分类", "YOLOv8 + 五路 CV 融合"),
    ("损伤几何尺寸\n量化", "像素-物理尺度换算"),
    ("三维空间定位", "解析式面投影"),
    ("检测报告\n自动生成", "python-docx"),
]
chevron_row(s, flow, Inches(0.4), Inches(1.8), 2.42, 1.5, gap=0.15,
            colors=[BLUE, BLUE, ORANGE, BLUE, DARK], title_size=14, note_size=10.5)
bar = add_rect(s, Inches(0.6), Inches(5.0), Inches(12.1), Inches(1.6), fill=LBLUE,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(bar.text_frame, [
    ("数据来源与模型策略", 14, DARK, True),
    ("公开数据集（GYU-DET / CODEBRIM）与预训练模型相结合，支持模型改进、数据增强与轻量化部署", 13, DARK, False),
], align=PP_ALIGN.CENTER, space_after=6)
for p in bar.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER

# ================= P7 技术架构 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "技术架构", "数据 — 算法 — 引擎 — 应用 四层架构 · PyTorch 开源框架训练推理", 7)
layers = [
    ("应用层", "桌面检测插件 · 三维可视化 · Word 报告自动生成", ORANGE),
    ("引擎层", "BIMBase 国产三维图形引擎 + pyp3d", DARK),
    ("算法层", "YOLOv8 目标检测 · 五路 CV 掩码融合 · Qwen-VL 多模态诊断", BLUE),
    ("数据层", "GYU-DET 11,123 张 / CODEBRIM 1,022 张 / 本地标注 1,014 张", GRAY),
]
for i, (t, d, col) in enumerate(layers):
    y = Inches(1.45) + i * Inches(1.3)
    txt_col = DARK if col is GRAY else WHITE
    add_rect(s, Inches(0.55), y, Inches(2.1), Inches(1.1), fill=col,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(0.55), y + Inches(0.33), Inches(2.1), Inches(0.5),
             [(t, 17, txt_col, True)], align=PP_ALIGN.CENTER)
    add_rect(s, Inches(2.85), y, Inches(6.9), Inches(1.1), fill=WHITE if col is GRAY else GRAY,
             line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(3.1), y + Inches(0.2), Inches(6.5), Inches(0.8),
             [(d, 12.5, DARK, False)])
# 层间箭头提示（用文字）
add_text(s, Inches(0.55), Inches(6.6), Inches(9.2), Inches(0.35),
         [("自下而上：数据 → 算法 → 引擎 → 应用（图中自下而上堆叠）", 10, DGRAY, False)],
         align=PP_ALIGN.CENTER)
ix = Inches(10.15)
fp = os.path.join(ROOT, "速构智维", "Picture", "桥隧识别.png")
if os.path.exists(fp):
    add_image_fit(s, fp, ix, Inches(1.6), Inches(2.6), Inches(2.6), 7, border=False)
cap = add_rect(s, ix, Inches(4.4), Inches(2.6), Inches(1.1), fill=LBLUE,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(cap.text_frame, [("自研插件图标", 12, DARK, True), ("桥隧识别", 10, DARK, False)],
         align=PP_ALIGN.CENTER, space_after=4)
cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# ================= P8 损伤识别与分类——方案 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "损伤智能识别与分类 —— 技术方案", "双模式检测：专用模型 + 免模型冷启动", 8)
# 两个模式卡片
m1 = add_rect(s, Inches(0.55), Inches(1.45), Inches(6.0), Inches(2.55), fill=WHITE, line=BLUE,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.85), Inches(1.6), Inches(5.4), Inches(2.3), [
    ("模式一：YOLOv8n 专用检测模型", 16, BLUE, True),
    ("· 覆盖 6 类表观损伤（裂缝 / 剥落 / 露筋 / 蜂窝麻面 / 渗水 / 锈胀）", 12.5, BODY, False),
    ("· 轻量化骨干网络，支持 CPU 端部署", 12.5, BODY, False),
    ("· 公开数据集 + 自建标注联合训练", 12.5, BODY, False),
], space_after=6)
m2 = add_rect(s, Inches(6.8), Inches(1.45), Inches(6.0), Inches(2.55), fill=WHITE, line=ORANGE,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(7.1), Inches(1.6), Inches(5.4), Inches(2.3), [
    ("模式二（特色）：免模型五路 CV 掩码融合", 16, ORANGE, True),
    ("· 黑帽变换 / 多尺度自适应阈值 / K-Means 颜色聚类", 12.5, BODY, False),
    ("· 局部纹理标准差 / Canny 边缘密度", 12.5, BODY, False),
    ("· HSV 天空抑制 + NMS 去重；无标注数据场景零门槛冷启动", 12.5, BODY, False),
], space_after=6)
# 数据集说明表
rows = [
    ("数据来源", "规模", "标注方式", "用途"),
    ("GYU-DET（公开）", "11,123 张", "已有标注（转换）", "训练 / 验证"),
    ("CODEBRIM（公开）", "1,022 张", "已有标注（转换）", "训练 / 验证"),
    ("本地巡检报告提取与人工标注", "1,014 张（整合集）", "YOLO 格式人工标注", "贴近真实业务场景"),
]
simple_table(s, rows, Inches(0.55), Inches(4.3), Inches(12.25), Inches(2.4),
             [Inches(4.6), Inches(2.4), Inches(3.0), Inches(2.25)])

# ================= P9 损伤识别与分类——指标与效果 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "损伤智能识别与分类 —— 指标与效果", "指标如实呈现，PR 曲线与混淆矩阵佐证准确率与召回率", 9)
# 左上指标卡
add_rect(s, Inches(0.55), Inches(1.35), Inches(6.0), Inches(1.75), fill=GRAY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.8), Inches(1.48), Inches(5.5), Inches(1.5), [
    ("YOLOv8n · 50 epochs · CPU 训练", 13, DARK, True),
    ("mAP50 ≈ 0.336 ｜ mAP50-95 ≈ 0.085", 18, ORANGE, True),
    ("（当前为基线模型，如实呈现；持续提升中）", 10.5, DGRAY, False),
], space_after=4)
add_rect(s, Inches(0.55), Inches(3.3), Inches(6.0), Inches(1.0), fill=WHITE, line=BLUE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.8), Inches(3.4), Inches(5.5), Inches(0.85), [
    ("提升路线：数据增强 · 更大骨干网络 · 半监督学习 · 类别均衡采样", 12, DARK, False),
])
add_image_fit(s, os.path.join(TRAIN_RUN, "results.png"),
              Inches(0.55), Inches(4.55), Inches(6.0), Inches(2.4), 9)
add_text(s, Inches(0.55), Inches(4.28), Inches(6.0), Inches(0.28),
         [("训练曲线（results.png）", 10.5, DARK, True)], align=PP_ALIGN.CENTER)
add_image_fit(s, os.path.join(VAL, "BoxPR_curve.png"),
              Inches(6.8), Inches(1.35), Inches(3.1), Inches(2.7), 9)
add_image_fit(s, os.path.join(VAL, "confusion_matrix_normalized.png"),
              Inches(10.05), Inches(1.35), Inches(3.0), Inches(2.7), 9)
add_image_fit(s, os.path.join(VAL, "val_batch0_pred.jpg"),
              Inches(6.8), Inches(4.55), Inches(6.25), Inches(2.4), 9)
add_text(s, Inches(6.8), Inches(4.28), Inches(6.25), Inches(0.28),
         [("典型损伤识别结果（val_batch0_pred.jpg）", 10.5, DARK, True)], align=PP_ALIGN.CENTER)
add_text(s, Inches(6.8), Inches(1.08), Inches(6.25), Inches(0.28),
         [("PR 曲线与归一化混淆矩阵：识别准确率 / 召回率佐证", 10.5, DARK, True)], align=PP_ALIGN.CENTER)

# ================= P10 损伤几何尺寸量化 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "损伤几何尺寸量化", "基于构件面投影标定的像素-物理尺度换算", 10)
bullets(s, Inches(0.55), Inches(1.4), Inches(7.5), Inches(1.7), [
    (0, "方法：由 2D→3D 面投影获得成像面物理尺寸，逐像素换算损伤长度 / 宽度 / 面积"),
], size=14)
qflow = [("检测框", ""), ("掩膜细化", ""), ("骨架提取", ""), ("长度/最大宽度/面积输出", "")]
chevron_row(s, qflow, Inches(0.55), Inches(2.35), 1.75, 0.85, gap=0.12,
            colors=[BLUE, BLUE, BLUE, ORANGE], title_size=12)
# 精度验证表（留占位）
rows = [
    ("验证项", "人工测量", "系统量化", "差异"),
    ("裂缝长度", "待补充实测数据", "待补充实测数据", "待补充实测数据"),
    ("损伤最大宽度", "待补充实测数据", "待补充实测数据", "待补充实测数据"),
    ("损伤面积", "待补充实测数据", "待补充实测数据", "待补充实测数据"),
]
simple_table(s, rows, Inches(0.55), Inches(3.6), Inches(7.5), Inches(2.2),
             [Inches(2.2), Inches(1.8), Inches(1.8), Inches(1.7)], size=12)
add_text(s, Inches(0.55), Inches(6.0), Inches(7.5), Inches(0.35),
         [("精度验证：与人工测量结果对比（测试完成后回填，原始记录可备查）", 10.5, DGRAY, False)])
placeholder(s, Inches(8.45), Inches(1.5), Inches(4.35), Inches(5.2),
            "尺寸量化结果展示截图（运行后替换）", 10)

# ================= P11 三维空间定位 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "三维空间定位", "解析式 2D→3D 面投影：轻量路线更适合基层推广", 11)
bullets(s, Inches(0.55), Inches(1.4), Inches(7.5), Inches(5.3), [
    (0, "方法：解析式 2D→3D 面投影"),
    (1, "构件面中心 + 法向 + U/V 轴解析公式等比投影"),
    (1, "2mm 法向偏移防闪面，点云标记 Combine 合成"),
    (0, "构件清单联动：病害归属到具体构件"),
    (1, "构件类型识别覆盖 19+ 种参数签名"),
    (0, "相比 SFM 三维重建 / GPS 定位："),
    (1, "无需额外硬件、计算量小、精度以构件为单元"),
    (1, "轻量路线更适合基层管养单位推广", ORANGE, True),
    (0, "定位精度说明：待补充实测数据"),
], size=14)
placeholder(s, Inches(8.45), Inches(1.5), Inches(4.35), Inches(5.2),
            "病害定位到 BIM 构件效果截图（BIMBase 内运行后替换）", 11)

# ================= P12 数字化成果形态 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "数字化成果形态", "可演示、可交互的数字化成果清单", 12)
cards = [
    ("BIMBase 桌面检测插件", "软件插件：检测→诊断→定位→报告全链路"),
    ("损伤识别交互界面", "照片导入、双模式检测、结果浏览"),
    ("三维可视化展示平台", "病害在 BIM 模型中三维标记展示"),
    ("检测报告自动生成工具", "python-docx 一键生成 Word 报告"),
]
for i, (t, d) in enumerate(cards):
    r, c = divmod(i, 2)
    x = Inches(0.6) + c * Inches(6.3)
    y = Inches(1.45) + r * Inches(1.85)
    add_rect(s, x, y, Inches(5.9), Inches(1.6), fill=GRAY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(1.6), fill=BLUE)
    add_text(s, x + Inches(0.35), y + Inches(0.2), Inches(5.2), Inches(1.3),
             [(f"{i+1}. {t}", 15.5, DARK, True), (d, 12, BODY, False)], space_after=6)
add_text(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(0.4),
         [("网页应用 / 小程序等成果形态已在路线图中规划", 12.5, ORANGE, True)], align=PP_ALIGN.CENTER)
placeholder(s, Inches(0.6), Inches(5.75), Inches(5.9), Inches(1.25),
            "检测插件运行界面截图（运行后替换）", 12)
placeholder(s, Inches(6.8), Inches(5.75), Inches(5.9), Inches(1.25),
            "三维可视化界面截图（运行后替换）", 12)

# ================= P13 检测报告自动生成 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "检测报告自动生成", "python-docx 一键生成 Word 检测报告，流程末环收口", 13)
rows = [
    ("章节", "内容"),
    ("封面信息", "工程名称 / 检测日期 / 检测单位 / 委托信息"),
    ("病害清单", "类型 · 位置（构件）· 尺寸 · 程度 · 紧迫度分级"),
    ("损伤位置分布图", "病害在桥梁三维模型 / 平面中的分布示意"),
    ("成因分析", "Qwen-VL 多模态诊断生成"),
    ("处置建议", "按病害类型与程度分级给出处置建议"),
]
simple_table(s, rows, Inches(0.55), Inches(1.5), Inches(7.5), Inches(4.2),
             [Inches(2.3), Inches(5.2)])
bullets(s, Inches(0.55), Inches(6.0), Inches(7.5), Inches(0.9), [
    (0, "报告规范、可归档，直接对接检测业务流程", ORANGE, True),
], size=13)
placeholder(s, Inches(8.45), Inches(1.5), Inches(4.35), Inches(5.2),
            "自动生成检测报告样例截图（Word 报告，运行后替换）", 13)

# ================= P14 案例验证与效果评价 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "案例验证与效果评价", "常泰长江大桥真实工程验证（BIM 模型 244.9MB · 15 类构件）", 14)
rows = [
    ("评价指标", "结果"),
    ("损伤识别准确率与召回率", "PR 曲线 / 混淆矩阵佐证（见 P9）"),
    ("单张图像平均检测时间", "待补充实测数据"),
    ("尺寸量化与人工测量差异", "待补充实测数据"),
    ("空间定位精度", "待补充实测数据"),
    ("人工判读环节减少情况", "判读、测量、定位、报告编制四环节：多人分步 → 一人一键流转（定性）"),
    ("检测报告生成效率改善", "一键自动生成规范 Word 报告，替代手工编制（定性）"),
]
simple_table(s, rows, Inches(0.55), Inches(1.45), Inches(12.25), Inches(4.7),
             [Inches(4.2), Inches(8.05)], size=12.5)
add_text(s, Inches(0.55), Inches(6.35), Inches(12.25), Inches(0.6),
         [("评价数据真实、清晰、可复核（测试脚本与原始记录可备查）；实测指标测试完成后回填", 12, ORANGE, True)],
         align=PP_ALIGN.CENTER)

# ================= P15 产教协同创新 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "产教协同创新", "企业出题 · 高校协同 · 学生攻关 · 成果服务产业", 15)
tri = [
    ("命题企业（广联达）", "业务需求 / 检测流程 / 数据格式", BLUE),
    ("高校教师", "土木工程 / 智能建造 / 计算机视觉专业指导", DARK),
    ("学生团队", "需求分析 / 方案设计 / 模型训练 / 系统开发 / 案例验证", ORANGE),
]
for i, (t, d, col) in enumerate(tri):
    x = Inches(0.6) + i * Inches(4.25)
    add_rect(s, x, Inches(1.5), Inches(4.0), Inches(2.1), fill=col,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.25), Inches(1.7), Inches(3.5), Inches(1.8),
             [(t, 16, WHITE, True), (d, 12, WHITE, False)], space_after=8)
    if i < 2:
        ar = add_rect(s, x + Inches(3.85), Inches(2.3), Inches(0.55), Inches(0.5),
                      fill=ORANGE, shape=MSO_SHAPE.LEFT_RIGHT_ARROW)
bullets(s, Inches(0.6), Inches(4.0), Inches(12.1), Inches(1.6), [
    (0, "跨专业组队建议：土木工程 + 计算机 + 人工智能 + 电子信息"),
    (0, "协同机制：企业提供真实业务命题与流程指导，高校教师把关技术路线，学生团队完成攻关与验证"),
], size=14)
add_rect(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.3), fill=GRAY, line=DGRAY,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.9), Inches(5.75), Inches(11.5), Inches(1.0),
         [("【过程记录占位】", 12, DGRAY, True),
          ("企业调研与技术交流过程说明（待补充）", 12, DGRAY, False)],
         align=PP_ALIGN.CENTER, space_after=4)
placed.append("P15: 企业调研与技术交流过程说明（待补充）")

# ================= P16 创新性与推广价值 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "创新性与推广价值", "四大核心创新 · 方向可拓展 · 方法可迁移", 16)
innov = [
    ("免模型五路 CV 融合", "黑帽 / 阈值 / 聚类 / 纹理 / 边缘五路掩码融合，零样本冷启动"),
    ("Qwen-VL 双图多模态诊断", "原图 + 带框图输入，五字段结构化 JSON 输出"),
    ("解析式面投影轻量定位", "替代昂贵 SFM / 点云路线，无额外硬件"),
    ("识别-量化-定位-报告闭环", "全链路数字化，一键贯通"),
]
for i, (t, d) in enumerate(innov):
    x = Inches(0.6) + i * Inches(3.15)
    add_rect(s, x, Inches(1.45), Inches(2.9), Inches(2.2), fill=WHITE, line=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.45), Inches(2.9), Inches(0.5), fill=DARK,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.12), Inches(1.53), Inches(2.66), Inches(0.38),
             [(f"{i+1}. {t}", 12.5, WHITE, True)])
    add_text(s, x + Inches(0.2), Inches(2.1), Inches(2.5), Inches(1.4),
             [(d, 11.5, BODY, False)])
bullets(s, Inches(0.6), Inches(3.95), Inches(12.1), Inches(2.2), [
    (0, "拓展方向（呼应命题鼓励内容）："),
    (1, "无人机 / 机器人平台边缘端轻量化部署 · 微裂缝等小目标检测优化"),
    (1, "多类型损伤同步识别 · 损伤时序对比与劣化趋势分析"),
    (1, "与桥梁管理系统 / 养护决策系统数据对接"),
    (0, "推广价值：方法可迁移至隧道、房屋等其他基础设施检测项目", ORANGE, True),
], size=13.5)

# ================= P17 成果提交物清单 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "成果提交物清单", "对照产业赛道答题要求逐项自查", 17)
rows = [
    ("序号", "提交物要求", "本项目对应", "状态"),
    ("①", "命题对策书 / 项目技术报告", "本 PPT + 技术文档", "✓"),
    ("②", "企业需求分析和总体解决方案", "P3-P7", "✓"),
    ("③", "业务流程图或技术路线图", "P6 五环流程 / P7 四层架构", "✓"),
    ("④", "可演示的数字化成果或交互原型", "BIMBase 检测插件", "✓"),
    ("⑤", "所选桥型与损伤类型的案例说明", "常泰长江大桥 · 6 类损伤", "✓"),
    ("⑥", "损伤识别、量化与定位的分析结果", "P8-P11", "✓"),
    ("⑦", "案例验证与应用效果说明", "P14", "✓"),
    ("⑧", "团队成员分工及校企协同过程说明", "待补充", "○"),
    ("⑨", "数据来源、知识产权及授权情况说明", "公开数据集 + 自建标注", "✓"),
]
simple_table(s, rows, Inches(0.55), Inches(1.4), Inches(12.25), Inches(5.4),
             [Inches(0.9), Inches(4.6), Inches(5.15), Inches(1.6)], size=12)

# ================= P18 应用边界与合规 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "应用边界与合规", "辅助决策工具 · 可解释可复核 · 数据合法合规", 18)
cards18 = [
    ("功能定位", "桥梁表观损伤智能检测的数字化辅助决策工具，不替代依法开展的定期检测、特殊检测、荷载试验和结构安全评估"),
    ("可解释可复核", "系统输出的损伤识别结果、尺寸数据和定位信息具备依据可查、过程可解释和人工可复核特征"),
    ("数据合规", "涉及实际工程图像、检测报告、企业数据及人员信息时，依法取得授权或进行脱敏处理"),
    ("知识产权", "成果真实合法，不侵犯他人知识产权；公开数据集按许可使用，自建标注归团队所有"),
]
for i, (t, d) in enumerate(cards18):
    r, c = divmod(i, 2)
    x = Inches(0.6) + c * Inches(6.3)
    y = Inches(1.55) + r * Inches(2.5)
    add_rect(s, x, y, Inches(5.9), Inches(2.2), fill=WHITE, line=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, y, Inches(0.12), Inches(2.2), fill=ORANGE)
    add_text(s, x + Inches(0.35), y + Inches(0.22), Inches(5.2), Inches(1.9),
             [(t, 17, DARK, True), (d, 12.5, BODY, False)], space_after=10)

# ================= P19 商业模式与市场前景 =================
s = prs.slides.add_slide(BLANK)
top_bar(s, "商业模式与市场前景", "全国 100 万+ 公路桥梁进入运维期，检测养护市场持续增长", 19)
biz = [
    ("软件授权", "检测插件 License 年费模式，面向检测单位与管养单位"),
    ("检测服务", "按项目收费：数据处理 + AI 分析 + 报告生成"),
    ("数据增值", "病害档案库建设与劣化趋势分析订阅服务"),
]
for i, (t, d) in enumerate(biz):
    x = Inches(0.6) + i * Inches(4.25)
    add_rect(s, x, Inches(1.5), Inches(4.0), Inches(1.9), fill=GRAY,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.5), Inches(4.0), Inches(0.55), fill=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x, Inches(1.6), Inches(4.0), Inches(0.4),
             [(t, 15, WHITE, True)], align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.25), Inches(2.2), Inches(3.5), Inches(1.1),
             [(d, 12, BODY, False)])
path = [("检测单位", "工具提效"), ("管养单位", "数字化档案"), ("桥隧全域", "基础设施推广")]
for i, (t, d) in enumerate(path):
    x = Inches(0.6) + i * Inches(4.25)
    ch = add_rect(s, x, Inches(3.8), Inches(3.9), Inches(1.1),
                  fill=[BLUE, DARK, ORANGE][i], shape=MSO_SHAPE.CHEVRON)
    tf = ch.text_frame; tf.word_wrap = True
    set_text(tf, [(t, 15, WHITE, True), (d, 11.5, WHITE, False)],
             align=PP_ALIGN.CENTER, space_after=2)
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
add_text(s, Inches(0.6), Inches(5.1), Inches(12.1), Inches(0.35),
         [("推广路径", 12, DGRAY, True)], align=PP_ALIGN.CENTER)
bar = add_rect(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.1), fill=LBLUE,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text(bar.text_frame, [
    ("落地门槛低", 14, DARK, True),
    ("普通无人机 + 笔记本即可部署，契合基层管养单位预算", 13, DARK, False),
], align=PP_ALIGN.CENTER, space_after=4)
for p in bar.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER

# ================= P20 结束页 =================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, fill=DARK)
add_rect(s, 0, Inches(4.6), W, Inches(0.05), fill=ORANGE)
add_text(s, Inches(1.0), Inches(1.9), Inches(11.3), Inches(1.1),
         [("让每一座桥梁的伤病，都被看见、被量化、被追踪", 36, WHITE, True)], align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(3.15), Inches(11.3), Inches(0.6),
         [("智检桥安 —— 服务桥梁全生命周期智慧运维", 17, LBLUE, False)],
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.2),
         [("感谢评审专家与命题企业（广联达）（待填写）", 14, WHITE, False),
          ("团队联系方式：________________（待填写）", 14, WHITE, False)],
         align=PP_ALIGN.CENTER, space_after=10)
n = len(thumbs_all)
show = [thumbs_all[i] for i in range(n) if i % max(1, n // 8) == 0][:8]
tw2 = Inches(1.15)
total = tw2 * len(show) + Inches(0.15) * (len(show) - 1)
x0 = int((W - total) / 2)
for i, (_, name, p) in enumerate(show):
    add_image_fit(s, p, Emu(int(x0 + i * (tw2 + Inches(0.15)))), Inches(5.95),
                  tw2, Inches(1.0), 20, border=False)

prs.save(OUT)
print("OK ->", OUT)
print("slides:", len(prs.slides._sldIdLst))
print("size:", os.path.getsize(OUT), "bytes")
print("\n--- 插入图片 ---")
for x in inserted:
    print(x)
print("\n--- 占位框 ---")
for x in placed:
    print(x)
